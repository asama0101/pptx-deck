from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER as T
from pptx.util import Emu

from pptx_deck import build, paths, template, theme
from pptx_deck.pptxutil import FONT

STANDARD = paths.ASSETS / "template" / "standard.potx"
PAL = theme.derive("#1F3A5F")


def run(model, tmp_path, tpl=STANDARD, footer=None):
    info = template.analyze(tpl)
    out = tmp_path / "o.pptx"
    build.build(model, tpl, info, PAL, out, base_dir=tmp_path, footer=footer, work_name="t")
    return Presentation(out), info


def text_el(idx, x, y, w, h, text, role="body", size=20, color="1A1A1A", bold=False, ph=None, **kw):
    return {"kind": "text", "idx": idx, "rect": {"x": x, "y": y, "w": w, "h": h}, "role": role, "ph": ph,
            "bg": "FFFFFF", "overlap": False, "chrome": False, "overflow": False,
            "box": {"x": x, "y": y, "w": w, "h": h},
            "paras": [{"align": kw.get("align", "left"), "bullet": kw.get("bullet", False), "lh": 1.4, "mb": 0,
                       "runs": [{"text": text, "size_px": size, "bold": bold, "italic": False, "color": color}]}]}


def slide(layout, elements, index=1):
    return {"index": index, "layout": layout, "bg": "FFFFFF", "elements": elements}


def shape_el(idx, x, y, w, h, kind="rect", fill="1F3A5F", a=1, radius=0, line=None):
    return {"kind": "shape", "idx": idx, "rect": {"x": x, "y": y, "w": w, "h": h}, "role": None, "ph": None,
            "bg": "FFFFFF", "overlap": False, "chrome": False, "shape": kind,
            "fill": {"hex": fill, "a": a} if fill else None, "line": line, "radius_px": radius}


def test_layouts_follow_data_layout(tmp_path):
    model = [slide("cover", [text_el(0, 96, 250, 900, 60, "表題", role=None, size=40, ph="title")], 1),
             slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28)], 2),
             slide("closing", [text_el(0, 96, 300, 900, 60, "以上", role=None, size=36, ph="title")], 3)]
    prs, info = run(model, tmp_path)
    assert [s.slide_layout.name for s in prs.slides] == ["表紙", "本文", "裏表紙"]
    assert [s.shapes.title.text_frame.text for s in prs.slides] == ["表題", "見出し", "以上"]


def test_user_template_slides_are_dropped_and_potx_footer_kept(tmp_path):
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28)])], tmp_path)
    s = prs.slides[0]
    types = [p.placeholder_format.type for p in s.placeholders]
    assert T.TITLE in types and T.SLIDE_NUMBER in types       # ページ番号はテンプレートのものを写す
    assert T.OBJECT not in types                               # 使わない本文領域は外す
    assert len(prs.slides) == 1


def test_missing_closing_uses_body(tmp_path):
    from tests.test_template import _as_pptx
    tpl = _as_pptx(STANDARD, tmp_path / "u.pptx", drop="裏表紙")
    model = [slide("closing", [text_el(0, 96, 300, 900, 60, "以上", role=None, size=36, ph="title")])]
    prs, info = run(model, tmp_path, tpl=tpl)
    assert info.closing is None and prs.slides[0].slide_layout.name == "本文"   # Review Focus 4


def test_title_placeholder_gets_measured_geometry_and_font(tmp_path):
    model = [slide("content", [text_el(0, 70, 44, 800, 40, "見出し", role="title", size=28, bold=True)])]
    prs, _ = run(model, tmp_path)
    t = prs.slides[0].shapes.title
    assert (t.left, t.top, t.width) == (70 * 9525, 44 * 9525, 800 * 9525)
    r = t.text_frame.paragraphs[0].runs[0]
    assert r.font.size.pt == 21 and r.font.bold and r.font.name == FONT
    assert r._r.rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea").get("typeface") == FONT


def test_textbox_bullets_and_runs(tmp_path):
    el = text_el(0, 80, 170, 360, 100, "x")
    el["paras"] = [
        {"align": "left", "bullet": True, "lh": 1.4, "mb": 6,
         "runs": [{"text": "太字", "size_px": 20, "bold": True, "italic": False, "color": "1F3A5F"},
                  {"br": True},
                  {"text": "続き", "size_px": 20, "bold": False, "italic": False, "color": "1A1A1A"}]}]
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), el])], tmp_path)
    tb = [s for s in prs.slides[0].shapes if not s.is_placeholder][0]
    assert tb.text_frame.word_wrap is True
    p = tb.text_frame.paragraphs[0]
    assert "buChar" in p._p.xml and [r.text for r in p.runs] == ["太字", "続き"]
    assert p.runs[0].font.bold is True


def test_shapes_fill_alpha_radius(tmp_path):
    els = [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28),
           shape_el(1, 64, 150, 300, 100, "round", radius=20),
           shape_el(2, 400, 150, 100, 100, "ellipse", fill="B85A0B", a=0.5),
           shape_el(3, 520, 150, 100, 100, "triangle", fill=None, line={"hex": "595959", "w": 2})]
    prs, _ = run([slide("content", els)], tmp_path)
    shapes = [s for s in prs.slides[0].shapes if not s.is_placeholder]
    assert shapes[0].auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert abs(shapes[0].adjustments[0] - 0.2) < 0.01
    assert shapes[1].auto_shape_type == MSO_SHAPE.OVAL and 'a:alpha val="50000"' in shapes[1]._element.xml.replace("'", '"')
    assert shapes[2].auto_shape_type == MSO_SHAPE.ISOSCELES_TRIANGLE and shapes[2].fill.type is not None
    # 影を付けない（LibreOffice はテーマの effectRef を見るため、参照も外す）
    assert all(s._element.xpath(".//a:effectRef/@idx") == ["0"] for s in shapes)


def test_line_arrow(tmp_path):
    ln = {"kind": "line", "idx": 1, "rect": {"x": 100, "y": 300, "w": 200, "h": 2}, "role": None, "ph": None,
          "bg": "FFFFFF", "overlap": False, "chrome": False, "fill": {"hex": "595959", "a": 1}, "arrow": True}
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), ln])], tmp_path)
    cx = [s for s in prs.slides[0].shapes if s.shape_type is not None and "Connector" in type(s).__name__][0]
    assert "tailEnd" in cx._element.xml and cx.begin_x == 100 * 9525 and cx.end_x == 300 * 9525


def test_line_arrow_reverse_swaps_ends(tmp_path):
    ln = {"kind": "line", "idx": 1, "rect": {"x": 100, "y": 300, "w": 200, "h": 2}, "role": None, "ph": None,
          "bg": "FFFFFF", "overlap": False, "chrome": False, "fill": {"hex": "595959", "a": 1}, "arrow": True,
          "reverse": True}
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), ln])], tmp_path)
    cx = [s for s in prs.slides[0].shapes if "Connector" in type(s).__name__][0]
    assert cx.begin_x == 300 * 9525 and cx.end_x == 100 * 9525          # 右から左へ向く


def test_icon_svg_becomes_picture_with_alt(tmp_path):
    svg = tmp_path / "i.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24">'
                   '<circle cx="12" cy="12" r="8" fill="none" stroke="#1F3A5F" stroke-width="2"/></svg>')
    im = {"kind": "image", "idx": 1, "rect": {"x": 100, "y": 200, "w": 48, "h": 48}, "role": "icon", "ph": None,
          "bg": "FFFFFF", "overlap": False, "chrome": False, "src": svg.as_uri(), "alt": "丸のアイコン"}
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), im])], tmp_path)
    pic = [s for s in prs.slides[0].shapes if s.shape_type == 13][0]
    assert pic._element.nvPicPr.cNvPr.get("descr") == "丸のアイコン"


def test_table_merge_fill_borders(tmp_path):
    def cell(c0, c1, text, fill=None, b=None, rowspan=1):
        return {"c0": c0, "c1": c1, "rowspan": rowspan, "fill": {"hex": fill, "a": 1} if fill else None,
                "borders": {"l": None, "r": None, "t": None, "b": b}, "pad": {"l": 8, "r": 8, "t": 8, "b": 8},
                "valign": "middle", "bg": "FFFFFF",
                "paras": [{"align": "left", "bullet": False, "lh": 1.4, "mb": 0,
                           "runs": [{"text": text, "size_px": 20, "bold": False, "italic": False, "color": "1A1A1A"}]}]}
    tbl = {"kind": "table", "idx": 1, "rect": {"x": 64, "y": 150, "w": 600, "h": 100}, "role": None, "ph": None,
           "bg": "FFFFFF", "overlap": False, "chrome": False,
           "table": {"xs": [64, 264, 464, 664], "rows": [
               {"h": 50, "cells": [cell(0, 1, "項目", fill="1F3A5F"), cell(1, 3, "見出し", b={"hex": "BFBFBF", "w": 2})]},
               {"h": 50, "cells": [cell(0, 1, "A"), cell(1, 2, "B"), cell(2, 3, "C")]}]}}
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), tbl])], tmp_path)
    gf = [s for s in prs.slides[0].shapes if s.has_table][0]
    t = gf.table
    assert len(t.columns) == 3 and len(t.rows) == 2
    assert t.cell(0, 1).is_merge_origin and t.cell(0, 1).span_width == 2
    assert t.cell(0, 0).text == "項目" and t.cell(1, 2).text == "C"
    assert "lnB" in t.cell(0, 1)._tc.xml


def test_footer_text_override(tmp_path):
    prs, _ = run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28)])],
                 tmp_path, footer="社外秘")
    ft = [p for p in prs.slides[0].placeholders if p.placeholder_format.type == T.FOOTER][0]
    assert ft.text_frame.text == "社外秘"


def test_chart_not_yet_supported_is_explicit(tmp_path):
    ch = {"kind": "chart", "idx": 1, "rect": {"x": 64, "y": 150, "w": 500, "h": 300}, "role": None, "ph": None,
          "bg": "FFFFFF", "overlap": False, "chrome": False,
          "chart": {"type": "bar", "categories": ["A"], "series": [{"name": "x", "values": [1]}]}}
    with pytest.raises(NotImplementedError):
        run([slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), ch])], tmp_path)
