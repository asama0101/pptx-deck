import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from pptx_deck import cli, initdir, paths
from tests.test_template import STANDARD, _as_pptx

pytestmark = pytest.mark.skipif(shutil.which("soffice") is None and not Path(
    r"C:\Program Files\LibreOffice\program\soffice.exe").exists(), reason="LibreOffice なし")

COMP = paths.ASSETS / "components"
ORDER = ["layouts/cover", "layouts/agenda", "layouts/section", "diagrams/flow-h", "tables/compare",
         "charts/bar", "diagrams/before-after", "diagrams/gantt", "layouts/decision", "layouts/closing"]


def make_deck(work: Path, names):
    slides = work / "slides"
    for i, n in enumerate(names, start=1):
        (slides / f"{i:02d}-A.html").write_text((COMP / f"{n}.html").read_text(encoding="utf-8"), encoding="utf-8")
    ans = "".join(f"## S{i:02d}\n- 選択: A\n- コメント: \n" for i in range(1, len(names) + 1))
    (work / "ans.md").write_text(ans, encoding="utf-8")
    assert cli.main(["variants", str(work)]) == 0
    assert cli.main(["finalize", str(work), "--answers", str(work / "ans.md")]) == 0


@pytest.mark.parametrize("template_kind", ["standard", "user_pptx"])
def test_full_deck_builds(tmp_path, capsys, template_kind):
    w = tmp_path / "deck"
    tpl = None
    if template_kind == "user_pptx":
        tpl = _as_pptx(STANDARD, tmp_path / "company.pptx",
                       rename={"表紙": "Title Slide", "本文": "Title and Content", "裏表紙": "Closing"})
    initdir.init(w, color="green", template_path=tpl)
    make_deck(w, ORDER)
    capsys.readouterr()
    assert cli.main(["build", str(w), "--footer", "社外秘"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["ok"] is True and out["violations"] == []
    prs = Presentation(w / "deck.pptx")
    assert len(prs.slides) == 10
    assert prs.slides[0].slide_layout.name in ("表紙", "Title Slide")
    assert prs.slides[-1].slide_layout.name in ("裏表紙", "Closing")
    assert any(s.has_chart for s in prs.slides[5].shapes) and any(s.has_table for s in prs.slides[4].shapes)
    footers = [p.text_frame.text for s in prs.slides for p in s.placeholders if "Footer" in p.name]
    assert footers and all(t == "社外秘" for t in footers)
    assert len(out["pngs"]) == 10


def test_no_picture_of_body_text(tmp_path, capsys):
    """本文を画像にしていない: 画像はアイコン等の小さなものだけ（SPEC 受入条件 4）。"""
    w = tmp_path / "deck"
    initdir.init(w, color="navy")
    make_deck(w, ["layouts/decision", "diagrams/compare", "charts/waterfall"])
    capsys.readouterr()
    assert cli.main(["build", str(w)]) == 0
    prs = Presentation(w / "deck.pptx")
    for s in prs.slides:
        for sh in s.shapes:
            if sh.shape_type == 13:
                assert sh.width * sh.height < prs.slide_width * prs.slide_height * 0.05
