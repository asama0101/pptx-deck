from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from pptx_deck import check, paths, template
from tests.test_build import shape_el, slide, text_el

INFO = template.analyze(paths.ASSETS / "template" / "standard.potx")
TITLE = text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28)
BOTTOM = text_el(90, 700, 560, 300, 60, "下端の本文")          # emptyBand を避けるため、領域の下端まで本文を置く


def rules(model):
    return sorted({v["rule"] for v in check.check_model(model, INFO)})


def test_clean_slide_passes():
    assert check.check_model([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "本文", size=20), BOTTOM])], INFO) == []


def test_overflow_flag():
    t = text_el(1, 64, 160, 500, 40, "本文", size=20)
    t["overflow"] = True
    assert "overflow" in rules([slide("content", [TITLE, t])])


def test_out_of_slide():
    assert "outOfSlide" in rules([slide("content", [TITLE, text_el(1, 1200, 160, 200, 40, "はみ出し")])])


def test_keepout_footer_zone():
    assert "keepout" in rules([slide("content", [TITLE, text_el(1, 64, 680, 500, 30, "フッター領域")])])


def test_chrome_elements_skip_keepout():
    c = text_el(1, 64, 680, 500, 30, "出典")
    c["chrome"] = True
    assert "keepout" not in rules([slide("content", [TITLE, c])])


def test_partial_text_overlap_is_error():
    a = text_el(1, 64, 160, 400, 60, "A")
    b = text_el(2, 300, 190, 400, 60, "B")
    assert "overlap" in rules([slide("content", [TITLE, a, b])])


def test_text_inside_shape_is_fine():
    sh = shape_el(1, 64, 150, 400, 120)
    t = text_el(2, 80, 170, 360, 40, "面の中", color="FFFFFF")
    t["bg"] = "1F3A5F"
    assert rules([slide("content", [TITLE, sh, t, BOTTOM])]) == []


def test_text_crossing_shape_edge_is_error():
    sh = shape_el(1, 64, 150, 200, 100)
    t = text_el(2, 200, 170, 300, 40, "はみ出す")
    assert "overlap" in rules([slide("content", [TITLE, sh, t])])


def test_data_overlap_allows_it():
    a = text_el(1, 64, 160, 400, 60, "A")
    b = text_el(2, 300, 190, 400, 60, "B")
    a["overlap"] = b["overlap"] = True
    assert "overlap" not in rules([slide("content", [TITLE, a, b])])


def test_font_min_by_role():
    assert "fontMin" in rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "小さい", size=16)])])   # 12pt
    assert "fontMin" not in rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "注記", role="caption", size=14)])])
    assert "fontMin" in rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "注記", role="caption", size=12)])])  # 9pt


def test_contrast_on_dark_template_background():          # Review Focus 3
    t = text_el(1, 64, 160, 500, 40, "暗い下地に黒文字", color="000000")
    t["bg"] = "101820"
    assert "contrast" in rules([slide("content", [TITLE, t])])


def test_large_text_needs_only_3_to_1():
    t = text_el(1, 64, 160, 500, 60, "大きい文字", size=36, color="8A8A8A")
    assert "contrast" not in rules([slide("content", [TITLE, t])])


def test_image_alt_and_body_image():
    im = {"kind": "image", "idx": 1, "rect": {"x": 64, "y": 150, "w": 900, "h": 450}, "role": None, "ph": None,
          "bg": "FFFFFF", "overlap": False, "chrome": False, "src": "x.png", "alt": ""}
    got = rules([slide("content", [TITLE, im])])
    assert "imageAlt" in got and "bodyImage" in got
    im2 = {**im, "role": "icon", "alt": "アイコン", "rect": {"x": 64, "y": 150, "w": 48, "h": 48}}
    assert rules([slide("content", [TITLE, im2, BOTTOM])]) == []


def test_title_missing_except_closing():
    assert "titleMissing" in rules([slide("content", [text_el(1, 64, 160, 500, 40, "本文")])])
    cover = text_el(0, 96, 250, 900, 60, "表題", role=None, size=40, ph="title")
    assert "titleMissing" not in rules([slide("cover", [cover])])


def test_check_pptx(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(12000000, 100000, 600000, 300000)     # 右にはみ出す
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "小さい"
    r.font.size = Pt(8)
    out = tmp_path / "x.pptx"
    prs.save(out)
    got = {v["rule"] for v in check.check_pptx(out)}
    assert {"pptxBounds", "pptxFontMin"} <= got


def test_errors_filter():
    vs = [{"rule": "a", "slide": 1, "message": "m", "severity": "error"},
          {"rule": "b", "slide": 1, "message": "m", "severity": "warning"}]
    assert [v["rule"] for v in check.errors(vs)] == ["a"]


def warn_rules(model):
    return sorted({v["rule"] for v in check.check_model(model, INFO) if v["severity"] == "warning"})


def test_dash_is_flagged_but_long_vowel_is_not():
    assert "dash" in warn_rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "A案——低い")])])
    assert "dash" not in warn_rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "サーバーとルーター")])])


def test_same_word_twice_is_flagged_short_values_are_not():
    a = text_el(1, 64, 160, 400, 40, "標準仕様データベースで代替")
    b = text_el(2, 64, 220, 400, 40, "標準仕様データベースで代替")
    assert "sameWord" in warn_rules([slide("content", [TITLE, a, b])])
    c, d = text_el(3, 64, 280, 100, 40, "38"), text_el(4, 200, 280, 100, 40, "38")
    assert "sameWord" not in warn_rules([slide("content", [TITLE, c, d])])


def test_orphan_last_line():
    t = text_el(1, 64, 160, 400, 80, "長い文が折り返して最後に二字")
    t["lastLineChars"] = 2
    assert "orphanLine" in warn_rules([slide("content", [TITLE, t])])
    t["lastLineChars"] = None
    assert "orphanLine" not in warn_rules([slide("content", [TITLE, t])])


def test_empty_band_when_body_stops_far_above_the_area_bottom():
    assert "emptyBand" in warn_rules([slide("content", [TITLE, text_el(1, 64, 160, 500, 40, "本文だけ")])])
    full = text_el(2, 64, 560, 500, 60, "下まで使う本文")
    assert "emptyBand" not in warn_rules([slide("content", [TITLE, full])])
    assert "emptyBand" not in warn_rules([slide("cover", [TITLE])])           # 表紙・章扉・目次・裏表紙は対象外
