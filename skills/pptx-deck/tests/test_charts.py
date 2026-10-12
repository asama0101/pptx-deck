import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE as C

from pptx_deck import charts, paths, template, theme
from pptx_deck.build import build
from tests.test_build import slide, text_el

STANDARD = paths.ASSETS / "template" / "standard.potx"
PAL = theme.derive("#1F3A5F")


def chart_el(spec):
    return {"kind": "chart", "idx": 1, "rect": {"x": 64, "y": 150, "w": 600, "h": 360}, "role": None, "ph": None,
            "bg": "FFFFFF", "overlap": False, "chrome": False, "chart": spec}


def built(spec, tmp_path):
    info = template.analyze(STANDARD)
    m = [slide("content", [text_el(0, 64, 40, 900, 40, "見出し", role="title", size=28), chart_el(spec)])]
    out = tmp_path / "c.pptx"
    build(m, STANDARD, info, PAL, out, base_dir=tmp_path)
    shapes = [s for s in Presentation(out).slides[0].shapes if s.has_chart]
    assert len(shapes) == 1
    return shapes[0].chart


BASE = {"categories": ["2024", "2025", "2026"], "series": [{"name": "費用", "values": [120, 90, 60]}]}


@pytest.mark.parametrize("t,expected", [
    ("bar", C.COLUMN_CLUSTERED), ("barh", C.BAR_CLUSTERED), ("line", C.LINE_MARKERS),
    ("stacked", C.COLUMN_STACKED), ("stacked100", C.COLUMN_STACKED_100)])
def test_chart_types(t, expected, tmp_path):
    spec = {**BASE, "type": t}
    if t.startswith("stacked"):
        spec["series"] = [{"name": "A", "values": [1, 2, 3]}, {"name": "B", "values": [3, 2, 1]}]
    ch = built(spec, tmp_path)
    assert ch.chart_type == expected


def test_pie(tmp_path):
    ch = built({"type": "pie", "categories": ["A", "B", "C"], "series": [{"name": "比率", "values": [50, 30, 20]}]}, tmp_path)
    assert ch.chart_type == C.PIE


def test_bar_highlight_uses_emphasis_color(tmp_path):
    ch = built({**BASE, "type": "bar", "highlight": 2}, tmp_path)
    pt = ch.plots[0].series[0].points[2]
    assert str(pt.format.fill.fore_color.rgb) == PAL.emphasis.lstrip("#")
    # 強調しない点は系列の塗り（基調色）のまま。点ごとの塗りは強調した点にだけ付く
    assert str(ch.plots[0].series[0].format.fill.fore_color.rgb) == PAL.accent.lstrip("#")


def test_fonts_are_at_least_10pt_and_biz(tmp_path):
    ch = built({**BASE, "type": "bar"}, tmp_path)
    assert ch.font.size.pt >= 10 and ch.font.name == "BIZ UDPゴシック"


def test_waterfall_series():
    d = charts.waterfall_series([100, 20, -30, 90], [0, 3])
    assert d["base"] == [0, 100, 90, 0] and d["up"] == [0, 20, 0, 0]
    assert d["down"] == [0, 0, 30, 0] and d["total"] == [100, 0, 0, 90]


def test_waterfall_default_totals_first_and_last():
    d = charts.waterfall_series([100, -40, 60], None)
    assert d["total"] == [100, 0, 60]


def test_waterfall_negative_running_total_is_an_error():
    with pytest.raises(ValueError, match="0 未満"):
        charts.waterfall_series([10, -30, 5], [0, 2])


def test_waterfall_chart_has_four_series_and_hidden_base(tmp_path):
    ch = built({"type": "waterfall", "categories": ["前", "増", "減", "後"],
                "series": [{"name": "金額", "values": [100, 20, -30, 90]}], "totals": [0, 3]}, tmp_path)
    s = ch.plots[0].series
    assert [x.name for x in s] == ["基準", "増加", "減少", "合計"] and len(s) == 4
    assert s[0].format.fill.type == 5          # MSO_FILL.BACKGROUND（塗りなし）


def test_waterfall_labels_are_readable_on_their_bars(tmp_path):
    ch = built({"type": "waterfall", "categories": ["前", "増", "減", "後"],
                "series": [{"name": "金額", "values": [100, 20, -30, 90]}], "totals": [0, 3]}, tmp_path)
    s = ch.plots[0].series
    fills = [PAL.accent, "#404040", "#767676"]
    for ser, fill in zip(list(s)[1:], fills):
        label = str(ser.data_labels.font.color.rgb)
        assert theme.contrast("#" + label, fill) >= 4.5


def test_unknown_type_is_an_error(tmp_path):
    with pytest.raises(ValueError, match="radar"):
        built({**BASE, "type": "radar"}, tmp_path)
