import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from pptx_deck import cli, initdir, variants

has_soffice = pytest.mark.skipif(shutil.which("soffice") is None and not Path(
    r"C:\Program Files\LibreOffice\program\soffice.exe").exists(), reason="LibreOffice なし")

COVER = """<section class="slide" data-layout="cover">
<div class="pp-text" data-ph="title">クラウド移行の承認依頼</div>
<div class="pp-text" data-ph="subtitle">2026年10月 情報システム部</div></section>"""

CONTENT = """<section class="slide" data-layout="content">
<div class="pp-text" data-role="title" style="height:44px">B案を採用し、年間5,800万円を削減する</div>
<div class="pp-shape" style="left:64px;top:160px;width:540px;height:300px;background:var(--accent-light)"></div>
<div class="pp-text" data-role="body" style="left:88px;top:184px;width:492px"><ul><li>初期費用は3案で最小</li><li>運用は既存チームで可能</li></ul></div>
<div class="pp-chart" data-chart='{"type":"bar","categories":["現行","B案"],"series":[{"name":"費用","values":[120,62]}],"highlight":1}' style="left:660px;top:160px;width:540px;height:300px"></div>
<div class="pp-text" data-role="caption" data-chrome style="left:64px;top:640px;width:700px;height:24px">出典: 社内試算（2026年9月）</div>
</section>"""


def prepare(tmp_path, content=CONTENT):
    w = tmp_path / "demo"
    initdir.init(w, color="navy")
    (w / "slides").mkdir(exist_ok=True)
    (w / "slides" / "01-A.html").write_text(COVER, encoding="utf-8")
    (w / "slides" / "02-A.html").write_text(content, encoding="utf-8")
    (w / "slides" / "02-B.html").write_text(content.replace("B案を採用", "B案を採る"), encoding="utf-8")
    return w


@has_soffice
def test_end_to_end(tmp_path, capsys):
    w = prepare(tmp_path)
    assert cli.main(["variants", str(w)]) == 0
    assert (w / "variants.html").exists()
    ans = tmp_path / "ans.md"
    ans.write_text("## S01\n- 選択: A\n- コメント: \n## S02\n- 選択: B\n- コメント: \n", encoding="utf-8")
    assert cli.main(["finalize", str(w), "--answers", str(ans)]) == 0
    capsys.readouterr()
    assert cli.main(["build", str(w)]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["ok"] is True and out["violations"] == []
    prs = Presentation(w / "demo.pptx")
    assert [s.slide_layout.name for s in prs.slides] == ["表紙", "本文"]
    assert prs.slides[1].shapes.title.text_frame.text.startswith("B案を採る")
    assert any(s.has_chart for s in prs.slides[1].shapes)
    assert len(out["pngs"]) == 2 and all(Path(p).exists() for p in out["pngs"])
    assert str(w) not in out["pngs"][0]                       # 確認用の画像は作業フォルダに出さない


def test_build_stops_on_errors_and_writes_no_pptx(tmp_path, capsys):
    bad = CONTENT.replace('style="left:88px;top:184px;width:492px"', 'style="left:88px;top:184px;width:492px;font-size:14px"')
    w = prepare(tmp_path, content=bad)
    ans = tmp_path / "ans.md"
    ans.write_text("## S01\n- 選択: A\n- コメント: \n## S02\n- 選択: A\n- コメント: \n", encoding="utf-8")
    cli.main(["finalize", str(w), "--answers", str(ans)])
    capsys.readouterr()
    assert cli.main(["build", str(w)]) == 1
    out = json.loads(capsys.readouterr().out)
    assert out["ok"] is False and any(v["rule"] == "fontMin" for v in out["violations"])
    assert not (w / "demo.pptx").exists()


def test_finalize_reports_all_problems(tmp_path, capsys):
    w = prepare(tmp_path)
    ans = tmp_path / "ans.md"
    ans.write_text("## S01\n- 選択: A\n- コメント: 直して\n", encoding="utf-8")
    assert cli.main(["finalize", str(w), "--answers", str(ans)]) == 2
    err = capsys.readouterr().err
    assert "S01" in err and "S02" in err


def test_finalize_merges_several_answer_files(tmp_path, capsys):
    w = prepare(tmp_path)
    a1, a2 = tmp_path / "a1.md", tmp_path / "a2.md"
    a1.write_text("## S01\n- 選択: A\n- コメント: \n## S02\n- 選択: B\n- コメント: 直して\n", encoding="utf-8")
    a2.write_text("## S01\n- 選択: \n- コメント: \n## S02\n- 選択: A\n- コメント: \n", encoding="utf-8")
    assert cli.main(["finalize", str(w), "--answers", str(a1)]) == 2          # コメント未反映で止まる
    capsys.readouterr()
    assert cli.main(["finalize", str(w), "--answers", str(a1), "--answers", str(a2)]) == 0
    assert (w / "final.html").exists()


def test_check_command(tmp_path, capsys):
    w = prepare(tmp_path)
    ans = tmp_path / "ans.md"
    ans.write_text("## S01\n- 選択: A\n- コメント: \n## S02\n- 選択: A\n- コメント: \n", encoding="utf-8")
    cli.main(["finalize", str(w), "--answers", str(ans)])
    capsys.readouterr()
    assert cli.main(["check", str(w / "final.html"), "--workdir", str(w)]) == 0
    assert json.loads(capsys.readouterr().out) == []
