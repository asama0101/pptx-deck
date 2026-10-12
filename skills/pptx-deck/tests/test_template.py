import io
import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from pptx_deck import paths, render, template
from pptx_deck.pptxutil import SLIDE_W_EMU

STANDARD = paths.ASSETS / "template" / "standard.potx"
has_soffice = pytest.mark.skipif(shutil.which("soffice") is None and not Path(
    r"C:\Program Files\LibreOffice\program\soffice.exe").exists(), reason="LibreOffice なし")


def _as_pptx(potx: Path, dst: Path, rename=None, drop=None) -> Path:
    """標準テンプレートから、利用者が持ってきそうな .pptx を作る（content-type を戻す）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(potx) as zin, zipfile.ZipFile(buf, "w") as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "[Content_Types].xml":
                d = d.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(it, d)
    buf.seek(0)
    prs = Presentation(buf)
    for lay in prs.slide_layouts:
        if rename and lay.name in rename:
            lay._element.cSld.set("name", rename[lay.name])
    if drop:
        for lay in list(prs.slide_layouts):
            if lay.name == drop:
                prs.slide_layouts.remove(lay)
    prs.save(dst)
    return dst


def test_standard_template_exists_and_opens_as_potx():
    assert STANDARD.exists()
    prs = template.open_template(STANDARD)  # .potx を開ける（Review Focus 1）
    assert [l.name for l in prs.slide_layouts] == ["表紙", "本文", "裏表紙"]


def test_analyze_standard():
    info = template.analyze(STANDARD)
    assert (info.cover, info.body, info.closing) == (0, 1, 2)
    assert info.area.to_dict() == {"x": 64, "y": 150, "w": 1152, "h": 490}
    assert info.title.y == 40
    assert any(r.y >= 670 for r in info.keepout)       # フッター・ページ番号の領域
    assert info.cover_subtitle is not None
    assert info.closing_title.y == 300                  # 裏表紙の題の位置（プレビューと PPTX をそろえる）
    assert not info.notes


def test_analyze_pptx_with_english_names(tmp_path):
    p = _as_pptx(STANDARD, tmp_path / "u.pptx",
                 rename={"表紙": "Title Slide", "本文": "Title and Content", "裏表紙": "Closing"})
    info = template.analyze(p)
    assert (info.cover, info.body, info.closing) == (0, 1, 2)


def test_missing_closing_falls_back_with_note(tmp_path):
    p = _as_pptx(STANDARD, tmp_path / "u.pptx", drop="裏表紙")
    info = template.analyze(p)
    assert info.closing is None                         # Review Focus 4
    assert any("裏表紙" in n for n in info.notes)
    assert info.closing_title == info.title             # 本文のレイアウトで代えるので、題は本文の位置


def test_non_169_template_is_rejected(tmp_path):
    prs = Presentation()  # 既定は 4:3（10in × 7.5in）
    src = tmp_path / "old.pptx"
    prs.save(src)
    with pytest.raises(ValueError, match="13.33"):      # Review Focus 2
        template.analyze(src)


def test_info_roundtrip():
    info = template.analyze(STANDARD)
    assert template.TemplateInfo.from_dict(info.to_dict()) == info


def test_layout_css_has_area_and_backgrounds():
    info = template.analyze(STANDARD)
    css = template.layout_css(info, "#FFFFFF")
    for token in ("--area-x: 64px", "--title-y: 40px", "--cover-title-x: 96px", "--closing-title-y: 300px",
                  "_template/bg-cover.png", "_template/bg-body.png", "--bg: #FFFFFF"):
        assert token in css


@has_soffice
def test_render_underlays(tmp_path):
    info = template.analyze(STANDARD)
    bg = template.render_underlays(STANDARD, info, tmp_path)
    for name in ("bg-cover.png", "bg-body.png", "bg-closing.png"):
        assert (tmp_path / name).exists()
    assert bg == "#FFFFFF"                              # 標準テンプレートの本文は白


@has_soffice
def test_pptx_to_pngs_uses_os_temp_and_returns_files(tmp_path):
    prs = template.open_template(STANDARD)
    prs.slides.add_slide(prs.slide_layouts[1])
    src = tmp_path / "x.pptx"
    prs.save(src)
    pngs = render.pptx_to_pngs(src, tmp_path / "out", width_px=640)
    assert len(pngs) == 1 and pngs[0].exists()
