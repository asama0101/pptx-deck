"""利用者（または標準）のテンプレートを開いて、使うレイアウトと領域を決める。"""
import io
import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as T
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .pptxutil import PX, SLIDE_H_EMU, SLIDE_W_EMU

_COVER = re.compile(r"(title slide|(?<!裏)表紙|cover|タイトル ?スライド)", re.I)
_CLOSING = re.compile(r"(裏表紙|closing|\bend\b|thank|結び)", re.I)
_BODY = re.compile(r"(title and content|本文|body|タイトルとコンテンツ|content)", re.I)


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    w: float
    h: float

    def to_dict(self) -> dict:
        return asdict(self)

    @staticmethod
    def from_dict(d: dict) -> "Rect":
        return Rect(**d)


@dataclass(frozen=True)
class TemplateInfo:
    cover: int
    body: int
    closing: int | None
    area: Rect
    title: Rect
    cover_title: Rect
    cover_subtitle: Rect | None
    closing_title: Rect
    keepout: tuple[Rect, ...] = ()
    notes: tuple[str, ...] = field(default_factory=tuple)

    def to_dict(self) -> dict:
        return {
            "cover": self.cover, "body": self.body, "closing": self.closing,
            "area": self.area.to_dict(), "title": self.title.to_dict(),
            "cover_title": self.cover_title.to_dict(),
            "cover_subtitle": self.cover_subtitle.to_dict() if self.cover_subtitle else None,
            "closing_title": self.closing_title.to_dict(),
            "keepout": [r.to_dict() for r in self.keepout], "notes": list(self.notes),
        }

    @staticmethod
    def from_dict(d: dict) -> "TemplateInfo":
        sub = d["cover_subtitle"]
        return TemplateInfo(
            cover=d["cover"], body=d["body"], closing=d["closing"],
            area=Rect.from_dict(d["area"]), title=Rect.from_dict(d["title"]),
            cover_title=Rect.from_dict(d["cover_title"]),
            cover_subtitle=Rect.from_dict(sub) if sub else None,
            closing_title=Rect.from_dict(d["closing_title"]),
            keepout=tuple(Rect.from_dict(r) for r in d["keepout"]), notes=tuple(d["notes"]))


def open_template(path) -> Presentation:
    """.pptx / .potx を開く。.potx は content-type を .pptx 相当に直して開く。"""
    path = Path(path)
    if path.suffix.lower() == ".potx":
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                zout.writestr(it, data)
        buf.seek(0)
        prs = Presentation(buf)
    else:
        prs = Presentation(str(path))
    if abs(prs.slide_width - SLIDE_W_EMU) > SLIDE_W_EMU * 0.01 or \
            abs(prs.slide_height - SLIDE_H_EMU) > SLIDE_H_EMU * 0.01:
        raise ValueError(
            "テンプレートは 13.33×7.5in（16:9）のみ対応です。"
            f"実際: {prs.slide_width / 914400:.2f}×{prs.slide_height / 914400:.2f}in")
    return prs


def _rect(shape) -> Rect:
    return Rect(shape.left / PX, shape.top / PX, shape.width / PX, shape.height / PX)


def _phs(layout, *types):
    return [p for p in layout.placeholders if p.placeholder_format.type in types]


def _pick(prs, name_re, pred=None):
    for i, lay in enumerate(prs.slide_layouts):
        if name_re.search(lay.name):
            return i
    if pred:
        for i, lay in enumerate(prs.slide_layouts):
            if pred(lay):
                return i
    return None


def _fixed_shapes(layout) -> tuple[Rect, ...]:
    """本文レイアウトとマスターにある、プレースホルダーでない図形（ロゴ・帯など）。全面の背景と面積 0 の線は除く。"""
    out = []
    for shapes in (layout.shapes, layout.slide_master.shapes):
        for sh in shapes:
            if sh.is_placeholder or None in (sh.left, sh.top, sh.width, sh.height):
                continue
            r = _rect(sh)
            if r.w * r.h <= 0 or r.w * r.h > 1280 * 720 * 0.5:
                continue
            out.append(r)
    return tuple(out)


def analyze(path) -> TemplateInfo:
    prs = open_template(path)
    L = list(prs.slide_layouts)
    notes = []
    body = _pick(prs, _BODY, lambda l: _phs(l, T.TITLE) and _phs(l, T.BODY, T.OBJECT))
    if body is None:
        raise ValueError("本文に使えるレイアウト（タイトルと本文の領域があるもの）がありません")
    cover = _pick(prs, _COVER, lambda l: bool(_phs(l, T.CENTER_TITLE)))
    if cover is None:
        cover = body
        notes.append("表紙のレイアウトが見つかりません。本文のレイアウトで代えます。")
    closing = _pick(prs, _CLOSING)
    if closing is None:
        notes.append("裏表紙のレイアウトが見つかりません。本文のレイアウトで代えます。")
    b, c = L[body], L[cover]
    cl_titles = _phs(L[closing], T.TITLE, T.CENTER_TITLE) if closing is not None else []
    areas = _phs(b, T.BODY, T.OBJECT)
    area = _rect(areas[0]) if areas else Rect(64, 150, 1152, 490)
    title = _rect(_phs(b, T.TITLE)[0])
    ct = _phs(c, T.CENTER_TITLE, T.TITLE)
    cs = _phs(c, T.SUBTITLE)
    keepout = tuple(_rect(p) for p in _phs(b, T.SLIDE_NUMBER, T.FOOTER, T.DATE)) + _fixed_shapes(b)
    return TemplateInfo(
        cover=cover, body=body, closing=closing, area=area, title=title,
        cover_title=_rect(ct[0]) if ct else title,
        cover_subtitle=_rect(cs[0]) if cs else None,
        closing_title=_rect(cl_titles[0]) if cl_titles else title,
        keepout=keepout, notes=tuple(notes))


_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
_THEME_KEYS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")


def _theme_colors(master) -> dict[str, str]:
    """テーマの色（dk1・lt1・accent1…）→ 6 桁の 16 進。"""
    root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    out = {}
    for key in _THEME_KEYS:
        el = root.find(f".//a:clrScheme/a:{key}", _NS)
        if el is None or len(el) == 0:
            continue
        c = el[0]
        val = c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr")
        if val:
            out[key] = val.upper()
    return out


def _color_hex(el, theme: dict[str, str], clrmap: dict[str, str]) -> str | None:
    if el.tag.endswith("srgbClr"):
        return el.get("val").upper()
    if el.tag.endswith("schemeClr"):
        key = el.get("val")
        return theme.get(clrmap.get(key, key))
    if el.tag.endswith("sysClr"):
        return (el.get("lastClr") or "FFFFFF").upper()
    return None


def layout_bg_hex(prs, layout_index: int) -> str | None:
    """レイアウト（なければマスター）の背景が単色なら '#RRGGBB'。画像・グラデーション・取れないときは None。"""
    layout = prs.slide_layouts[layout_index]
    master = layout.slide_master
    theme = _theme_colors(master)
    cm = master._element.find("p:clrMap", _NS)
    clrmap = dict(cm.attrib) if cm is not None else {}
    for part in (layout._element, master._element):
        bg = part.find("p:cSld/p:bg", _NS)
        if bg is None:
            continue
        solid = bg.find("p:bgPr/a:solidFill", _NS)
        if solid is not None and len(solid):
            h = _color_hex(solid[0], theme, clrmap)
            return f"#{h}" if h else None
        ref = bg.find("p:bgRef", _NS)
        if ref is not None and len(ref):                 # テーマの背景スタイルの参照。色の指定がある（bg1 など）
            h = _color_hex(ref[0], theme, clrmap)
            return f"#{h}" if h else None
        return None                                      # 画像・グラデーションなど
    return None


def background_colors(path, info: TemplateInfo) -> tuple[dict[str, str], list[str]]:
    """表紙・本文・裏表紙の背景色（HTML の下地）。単色でないレイアウトは白で近似し、notes に理由を返す。"""
    prs = open_template(path)
    idx = {"cover": info.cover, "body": info.body, "closing": info.closing if info.closing is not None else info.body}
    out, notes = {}, []
    for name, i in idx.items():
        h = layout_bg_hex(prs, i)
        if h is None:
            h = "#FFFFFF"
            notes.append(f"{name} の背景が単色ではないため、HTML の下地は白で近似しました（PPTX にはテンプレートの背景が出ます）。")
        out[name] = h
    return out, notes


def layout_css(info: TemplateInfo, bg_hex: str, layout_bg: dict[str, str] | None = None) -> str:
    def v(prefix, r: Rect):
        return (f"  --{prefix}-x: {r.x:g}px;\n  --{prefix}-y: {r.y:g}px;\n"
                f"  --{prefix}-w: {r.w:g}px;\n  --{prefix}-h: {r.h:g}px;\n")
    root = f":root {{\n  --bg: {bg_hex};\n" + v("area", info.area) + v("title", info.title) \
        + v("cover-title", info.cover_title) + v("closing-title", info.closing_title)
    if info.cover_subtitle:
        root += v("cover-subtitle", info.cover_subtitle)
    root += "}\n"
    for name in ("cover", "closing"):       # 表紙・裏表紙の下地が本文と違う色でも、文字色の検査が合うようにする
        if layout_bg and layout_bg.get(name) and layout_bg[name] != bg_hex:
            root += f'.slide[data-layout="{name}"] {{ --bg: {layout_bg[name]}; }}\n'
    return root
