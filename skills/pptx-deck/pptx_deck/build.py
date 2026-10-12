"""モデル → PPTX。テンプレートのファイルを土台に、レイアウトの上へ書き出す。"""
import copy
import re
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import url2pathname

from lxml import etree
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER as T
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import paths, template
from .pptxutil import drop_all_slides, emu, rgb, set_run_font

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
          "ellipse": MSO_SHAPE.OVAL, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def url_to_path(url: str, base_dir: Path) -> Path:
    if url.startswith("file:"):
        return Path(url2pathname(urlparse(url).path))
    return (base_dir / url).resolve()


# ---- 文字 ----
def _bullet(p, size_px):
    pPr = p._p.get_or_add_pPr()
    indent = round(size_px * 9525)
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def fill_text_frame(tf, paras, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[para["align"]]
        size = next((r["size_px"] for r in para["runs"] if "text" in r), 20)
        p.line_spacing = max(0.9, round(para["lh"] / 1.2, 2))
        if para.get("mb"):
            p.space_after = Pt(round(para["mb"] * 0.75, 1))
        if para["bullet"]:
            _bullet(p, size)
        for run in para["runs"]:
            if run.get("br"):
                p.add_line_break()
                continue
            r = p.add_run()
            r.text = run["text"]
            set_run_font(r, run["size_px"], run["bold"], run["italic"], run["color"])


def _place(shape, rect):
    shape.left, shape.top = emu(rect["x"]), emu(rect["y"])
    shape.width, shape.height = emu(rect["w"]), emu(rect["h"])


def add_textbox(slide, el):
    b = el["box"]
    tb = slide.shapes.add_textbox(emu(b["x"]), emu(b["y"]), emu(b["w"] + 2), emu(max(b["h"], 1)))
    fill_text_frame(tb.text_frame, el["paras"])


# ---- 図形・線・画像 ----
def add_shape(slide, el):
    r = el["rect"]
    shp = slide.shapes.add_shape(SHAPES[el["shape"]], emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    if el["shape"] == "round" and el["radius_px"]:
        shp.adjustments[0] = min(0.5, el["radius_px"] / min(r["w"], r["h"]))
    if el["fill"]:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(el["fill"]["hex"])
        if el["fill"]["a"] < 1:
            clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(el["fill"]["a"] * 100000)))
    else:
        shp.fill.background()
    if el["line"]:
        shp.line.color.rgb = rgb(el["line"]["hex"])
        shp.line.width = Pt(el["line"]["w"] * 0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    for ref in shp._element.xpath(".//a:effectRef"):  # LibreOffice は effectLst を無視してテーマの影を出す
        ref.set("idx", "0")


def add_line(slide, el):
    r = el["rect"]
    horizontal = r["w"] >= r["h"]
    if horizontal:
        x1, y1, x2, y2, th = r["x"], r["y"] + r["h"] / 2, r["x"] + r["w"], r["y"] + r["h"] / 2, r["h"]
    else:
        x1, y1, x2, y2, th = r["x"] + r["w"] / 2, r["y"], r["x"] + r["w"] / 2, r["y"] + r["h"], r["w"]
    if el.get("reverse"):
        x1, y1, x2, y2 = x2, y2, x1, y1
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    c.line.color.rgb = rgb(el["fill"]["hex"])
    c.line.width = Pt(max(0.75, th * 0.75))
    if el["arrow"]:
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")


def rasterize_svg(svg: Path, width_px: float, work_name: str) -> Path:
    import fitz
    out_dir = paths.temp_dir(work_name) / "raster"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (svg.stem + ".png")
    doc = fitz.open(str(svg))
    page = doc[0]
    zoom = max(1.0, width_px * 3 / page.rect.width)
    page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True).save(str(out))
    doc.close()
    return out


def add_image(slide, el, base_dir, work_name):
    r = el["rect"]
    src = url_to_path(el["src"], base_dir)
    if src.suffix.lower() == ".svg":
        src = rasterize_svg(src, r["w"], work_name)
    pic = slide.shapes.add_picture(str(src), emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    pic._element.nvPicPr.cNvPr.set("descr", el["alt"])


# ---- 表 ----
def _border(tcPr, tag, spec):
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    ln = etree.Element(qn(tag))
    if spec:
        ln.set("w", str(int(spec["w"] * 0.75 * 12700)))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", spec["hex"])
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    return ln


def add_table(slide, el):
    t, r = el["table"], el["rect"]
    xs, rows = t["xs"], t["rows"]
    ncols = len(xs) - 1
    gf = slide.shapes.add_table(len(rows), ncols, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for i in range(ncols):
        tbl.columns[i].width = emu(xs[i + 1] - xs[i])
    for j, row in enumerate(rows):
        tbl.rows[j].height = emu(row["h"])
    for j, row in enumerate(rows):
        for c in row["cells"]:
            cell = tbl.cell(j, c["c0"])
            if c["c1"] - c["c0"] > 1 or c["rowspan"] > 1:
                cell.merge(tbl.cell(j + c["rowspan"] - 1, c["c1"] - 1))
            cell.margin_left, cell.margin_right = emu(c["pad"]["l"]), emu(c["pad"]["r"])
            cell.margin_top, cell.margin_bottom = emu(c["pad"]["t"]), emu(c["pad"]["b"])
            cell.vertical_anchor = VALIGN[c["valign"]]
            if c["fill"]:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(c["fill"]["hex"])
            else:
                cell.fill.background()
            fill_text_frame(cell.text_frame, c["paras"], anchor=VALIGN[c["valign"]])
            cell.margin_left, cell.margin_right = emu(c["pad"]["l"]), emu(c["pad"]["r"])
            cell.margin_top, cell.margin_bottom = emu(c["pad"]["t"]), emu(c["pad"]["b"])
            tcPr = cell._tc.get_or_add_tcPr()
            for i, (tag, side) in enumerate((("a:lnL", "l"), ("a:lnR", "r"), ("a:lnT", "t"), ("a:lnB", "b"))):
                tcPr.insert(i, _border(tcPr, tag, c["borders"][side]))


# ---- スライド ----
def _next_id(slide) -> int:
    ids = [int(x) for x in slide.shapes._spTree.xpath("//p:cNvPr/@id")]
    return max(ids + [1]) + 1


def _clone_footer_placeholders(slide, layout, footer):
    for ph in layout.placeholders:
        if ph.placeholder_format.type not in (T.SLIDE_NUMBER, T.FOOTER):
            continue
        el = copy.deepcopy(ph._element)
        el.xpath(".//p:cNvPr")[0].set("id", str(_next_id(slide)))
        slide.shapes._spTree.append(el)
        if footer is not None and ph.placeholder_format.type == T.FOOTER:
            ts = el.xpath(".//a:t")
            if ts:
                ts[0].text = footer
                for extra in ts[1:]:
                    extra.text = ""
            else:
                p = el.xpath(".//a:p")[0]
                r = etree.Element(qn("a:r"))
                etree.SubElement(r, qn("a:t")).text = footer
                end = p.find(qn("a:endParaRPr"))
                if end is not None:      # endParaRPr は段落の最後に置く
                    end.addprevious(r)
                else:
                    p.append(r)


def _layout_index(info, name):
    if name == "cover":
        return info.cover
    if name == "closing":
        return info.closing if info.closing is not None else info.body
    return info.body


def build(model, template_path, info, palette, out: Path, base_dir: Path,
          footer: str | None = None, work_name: str = "build") -> Path:
    from . import charts

    prs = template.open_template(template_path)
    drop_all_slides(prs)
    for s in model:
        layout = prs.slide_layouts[_layout_index(info, s["layout"])]
        slide = prs.slides.add_slide(layout)
        used = set()
        elements = s["elements"]
        # タイトル・副題はレイアウトのプレースホルダーへ
        for el in elements:
            if el["kind"] != "text":
                continue
            is_title = el.get("ph") == "title" or (el.get("role") == "title" and s["layout"] not in ("cover", "closing"))
            is_sub = el.get("ph") == "subtitle"
            if is_title and slide.shapes.title is not None and "title" not in used:
                ph = slide.shapes.title
                if el.get("ph") is None:        # 本文のタイトルは、測った位置に上書きする
                    _place(ph, el["box"])
                fill_text_frame(ph.text_frame, el["paras"])
                used.add("title")
                el["_done"] = True
            elif is_sub:
                subs = [p for p in slide.placeholders if p.placeholder_format.type == T.SUBTITLE]
                if subs:
                    fill_text_frame(subs[0].text_frame, el["paras"])
                    used.add("sub")
                    el["_done"] = True
        for ph in list(slide.placeholders):
            keep = (ph.placeholder_format.type in (T.TITLE, T.CENTER_TITLE) and "title" in used) or \
                   (ph.placeholder_format.type == T.SUBTITLE and "sub" in used)
            if not keep:
                ph._element.getparent().remove(ph._element)
        _clone_footer_placeholders(slide, layout, footer)
        for el in elements:
            if el.pop("_done", False):
                continue
            k = el["kind"]
            if k == "text":
                add_textbox(slide, el)
            elif k == "shape":
                add_shape(slide, el)
            elif k == "line":
                add_line(slide, el)
            elif k == "image":
                add_image(slide, el, base_dir, work_name)
            elif k == "table":
                add_table(slide, el)
            elif k == "chart":
                charts.add_chart(slide, el, palette)
            else:
                raise ValueError(f"未対応の要素: {k}")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
