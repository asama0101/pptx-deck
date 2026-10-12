"""体裁の検査（最小限）: はみ出し・重なり・文字サイズの下限・コントラスト・画像の使い方。"""
from pathlib import Path

from . import theme
from .template import TemplateInfo

SLIDE_W, SLIDE_H = 1280, 720
MIN_PT_DEFAULT, MIN_PT_SMALL = 14, 10          # 本文 / 注記（role=caption|chrome）
LARGE_PX = 24                                  # 24px(18pt) 以上は大きい文字（3:1）
BODY_IMAGE_RATIO = 0.4
OVERLAP_MIN_AREA = 4


def _v(rule, slide, message, severity="error"):
    return {"rule": rule, "slide": slide, "message": message, "severity": severity}


def errors(violations):
    return [v for v in violations if v["severity"] == "error"]


def _r(el):
    return el.get("box") if el["kind"] == "text" else el["rect"]


def _inter(a, b):
    ix = min(a["x"] + a["w"], b["x"] + b["w"]) - max(a["x"], b["x"])
    iy = min(a["y"] + a["h"], b["y"] + b["h"]) - max(a["y"], b["y"])
    return ix * iy if ix > 0 and iy > 0 else 0


def _contains(a, b, tol=2):
    return (a["x"] - tol <= b["x"] and a["y"] - tol <= b["y"]
            and a["x"] + a["w"] + tol >= b["x"] + b["w"] and a["y"] + a["h"] + tol >= b["y"] + b["h"])


def check_model(model, info: TemplateInfo):
    out = []
    keepout = [r.to_dict() for r in info.keepout]
    for s in model:
        n = s["index"]
        els = s["elements"]
        has_title = any(e["kind"] == "text" and (e.get("role") == "title" or e.get("ph") == "title") for e in els)
        if not has_title and s["layout"] != "closing":
            out.append(_v("titleMissing", n, "タイトル（data-role=\"title\" または data-ph=\"title\"）がありません"))
        for e in els:
            k, rect = e["kind"], _r(e)
            label = f"{k}#{e['idx']}"
            if k == "text":
                if e.get("overflow"):
                    out.append(_v("overflow", n, f"{label}: 文字が枠からはみ出しています。幅・高さを広げるか文章を減らしてください"))
                for p in e["paras"]:
                    for run in p["runs"]:
                        if "text" not in run:
                            continue
                        pt = run["size_px"] * 0.75
                        small = e.get("role") in ("caption", "chrome") or e.get("chrome")
                        floor = MIN_PT_SMALL if small else MIN_PT_DEFAULT
                        if pt < floor:
                            out.append(_v("fontMin", n, f"{label}: {pt:g}pt は下限 {floor}pt 未満です"))
                        need = 3.0 if run["size_px"] >= LARGE_PX else 4.5
                        ratio = theme.contrast("#" + run["color"], "#" + e["bg"])
                        if ratio < need:
                            out.append(_v("contrast", n, f"{label}: 文字色 #{run['color']} と背景 #{e['bg']} のコントラスト比 {ratio:.1f} が {need} 未満です"))
            if k == "table":
                for row in e["table"]["rows"]:
                    for c in row["cells"]:
                        for p in c["paras"]:
                            for run in p["runs"]:
                                if "text" in run and run["size_px"] * 0.75 < MIN_PT_DEFAULT:
                                    out.append(_v("fontMin", n, f"{label}: 表の文字 {run['size_px'] * 0.75:g}pt は下限 {MIN_PT_DEFAULT}pt 未満です"))
                                bg = c["fill"]["hex"] if c["fill"] else c["bg"]
                                need = 3.0 if run.get("size_px", 0) >= LARGE_PX else 4.5
                                if "text" in run and theme.contrast("#" + run["color"], "#" + bg) < need:
                                    out.append(_v("contrast", n, f"{label}: 表の文字色 #{run['color']} と背景 #{bg} のコントラスト比が {need} 未満です"))
            if k == "image":
                if not e.get("alt"):
                    out.append(_v("imageAlt", n, f"{label}: alt がありません"))
                area = rect["w"] * rect["h"]
                if e.get("role") not in ("icon", "logo", "photo") or area > SLIDE_W * SLIDE_H * BODY_IMAGE_RATIO:
                    out.append(_v("bodyImage", n, f"{label}: 本文を画像にしないでください（画像はアイコン・ロゴ・利用者の写真だけ）"))
            if rect["x"] < -1 or rect["y"] < -1 or rect["x"] + rect["w"] > SLIDE_W + 1 or rect["y"] + rect["h"] > SLIDE_H + 1:
                out.append(_v("outOfSlide", n, f"{label}: スライドの外にはみ出しています"))
            if not e.get("chrome") and not e.get("ph") and k != "line":
                for z in keepout:
                    if _inter(rect, z) > OVERLAP_MIN_AREA:
                        out.append(_v("keepout", n, f"{label}: テンプレートのヘッダー・フッター領域に入っています"))
                        break
        # 重なり
        cand = [e for e in els if e["kind"] in ("text", "shape", "image", "table", "chart")
                and not (e["kind"] == "text" and not e["paras"])]
        for i, a in enumerate(cand):
            for b in cand[i + 1:]:
                if a["overlap"] or b["overlap"]:
                    continue
                ra, rb = _r(a), _r(b)
                if _inter(ra, rb) <= OVERLAP_MIN_AREA:
                    continue
                if (a["kind"] == "shape" and _contains(ra, rb)) or (b["kind"] == "shape" and _contains(rb, ra)):
                    continue
                out.append(_v("overlap", n, f"{a['kind']}#{a['idx']} と {b['kind']}#{b['idx']} が重なっています（意図的なら data-overlap を両方に付ける）"))
    return out


def check_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    tol = 9525
    out = []
    for i, s in enumerate(prs.slides, start=1):
        for sh in s.shapes:
            if sh.left is None:
                continue
            if sh.left < -tol or sh.top < -tol or sh.left + sh.width > W + tol or sh.top + sh.height > H + tol:
                out.append(_v("pptxBounds", i, f"{sh.name}: スライドの外にはみ出しています"))
            if sh.shape_type == 13 and sh.width * sh.height > W * H * BODY_IMAGE_RATIO:
                out.append(_v("pptxPicture", i, f"{sh.name}: 大きな画像があります（本文を画像にしない）"))
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.font.size.pt < MIN_PT_SMALL:
                            out.append(_v("pptxFontMin", i, f"{sh.name}: {r.font.size.pt:g}pt は下限 {MIN_PT_SMALL}pt 未満です"))
    return out
