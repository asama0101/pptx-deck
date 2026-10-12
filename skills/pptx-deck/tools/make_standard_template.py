"""標準テンプレート assets/template/standard.potx を生成する。再生成は:
   python -m tools.make_standard_template  （venv の Python、skills/pptx-deck で実行）
レイアウトは 表紙・本文・裏表紙 の 3 つ。装飾は無彩色の細線だけ（どの配色にも合う）。"""
import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as T
from pptx.oxml import parse_xml

PX = 9525
OUT = Path(__file__).resolve().parent.parent / "assets" / "template" / "standard.potx"
NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')

FOOT = {T.DATE: (800, 676, 300, 28), T.FOOTER: (64, 676, 700, 28), T.SLIDE_NUMBER: (1136, 676, 80, 28)}
GEOM = {
    "表紙": {T.CENTER_TITLE: (96, 250, 1088, 120), T.SUBTITLE: (96, 390, 1088, 60), **FOOT},
    "本文": {T.TITLE: (64, 40, 1152, 80), T.OBJECT: (64, 150, 1152, 490), **FOOT},
    "裏表紙": {T.TITLE: (96, 300, 1088, 100), **FOOT},
}
# レイアウト名 → (線の始点 x, y, 長さ)
LINES = {"表紙": (96, 236, 120), "本文": (64, 668, 1152), "裏表紙": (96, 286, 120)}


def _line_xml(shape_id, x, y, length):
    return parse_xml(
        f'<p:cxnSp {NS}><p:nvCxnSpPr><p:cNvPr id="{shape_id}" name="罫線"/><p:cNvCxnSpPr/>'
        f'<p:nvPr userDrawn="1"/></p:nvCxnSpPr><p:spPr><a:xfrm><a:off x="{x * PX}" y="{y * PX}"/>'
        f'<a:ext cx="{length * PX}" cy="0"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
        f'<a:ln w="9525"><a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill></a:ln></p:spPr></p:cxnSp>')


def _place(shape, geom):
    x, y, w, h = geom
    shape.left, shape.top, shape.width, shape.height = (round(v * PX) for v in (x, y, w, h))


def make(out: Path = OUT) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    for ph in prs.slide_master.placeholders:
        t = ph.placeholder_format.type
        g = {T.TITLE: (64, 40, 1152, 80), T.BODY: (64, 150, 1152, 490), **FOOT}.get(t)
        if g:
            _place(ph, g)
    layouts = list(prs.slide_layouts)
    keep = {0: "表紙", 1: "本文", 5: "裏表紙"}  # 既定テンプレートの Title Slide / Title and Content / Title Only
    kept = {name: layouts[i] for i, name in keep.items()}
    for i, lay in enumerate(layouts):
        if i not in keep:
            prs.slide_layouts.remove(lay)
    for name, lay in kept.items():
        lay._element.cSld.set("name", name)
        for ph in lay.placeholders:
            g = GEOM[name].get(ph.placeholder_format.type)
            if g:
                _place(ph, g)
        x, y, length = LINES[name]
        lay.shapes._spTree.append(_line_xml(900, x, y, length))
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml",
                                    b"presentationml.template.main+xml")
            zout.writestr(item, data)
    return out


if __name__ == "__main__":
    print(make())
