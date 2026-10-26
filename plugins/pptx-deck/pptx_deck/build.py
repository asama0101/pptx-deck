"""モデル → PPTX。テンプレートのファイルを土台に、レイアウトの上へ書き出す。"""

import hashlib
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER as T
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import paths, template
from .pptxutil import clone_footer_placeholders, drop_all_slides, emu, file_url_to_path, rgb, set_run_font

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
PICTURE_NAME = "pp-picture"  # 絵（pp-picture）として貼った画像の名前。check.check_pptx が大きな画像の検査から除く


def url_to_path(url: str, base_dir: Path) -> Path:
    if url.startswith("file:"):
        return file_url_to_path(url)
    return (base_dir / url).resolve()


# ---- 文字 ----
def _bullet(p, size_px):
    pPr = p._p.get_or_add_pPr()
    indent = round(size_px * 9525)
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def fill_text_frame(tf, paras, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[para["align"]]
        size = next((r["size_px"] for r in para["runs"] if "text" in r), 20)
        # 固定ポイント。倍率だと、フォントの行の高さで間隔が変わる。行は最大の文字に合わせて伸びる
        line_px = max((r["size_px"] for r in para["runs"] if "text" in r), default=size)
        p.line_spacing = Pt(round(para["lh"] * line_px * 0.75, 2))
        if para.get("mb"):
            p.space_after = Pt(round(para["mb"] * 0.75, 1))
        if para["bullet"]:
            _bullet(p, size)
        for run in para["runs"]:
            if run.get("br"):
                p.add_line_break()
                continue
            r = p.add_run()
            r.text = run["text"]
            set_run_font(r, run["size_px"], run["bold"], run["italic"], run["color"])


def _place(shape, rect):
    shape.left, shape.top = emu(rect["x"]), emu(rect["y"])
    shape.width, shape.height = emu(rect["w"]), emu(rect["h"])


def add_textbox(slide, el):
    b = el["box"]
    width = min(b["w"] + 2, 1280 - b["x"])  # 折り返しの誤差で +2px。ただしスライドの右端は越えない
    tb = slide.shapes.add_textbox(emu(b["x"]), emu(b["y"]), emu(width), emu(max(b["h"], 1)))
    fill_text_frame(tb.text_frame, el["paras"])


# ---- 図形・線・画像 ----
def _set_dash(line_fmt, dash):
    """a:prstDash を a:ln の中のスキーマ順（塗りの後ろ、round/headEnd/tailEnd の前）に置く。"""
    if not dash:
        return
    ln = line_fmt._get_or_add_ln()
    for old in ln.findall(qn("a:prstDash")):
        ln.remove(old)
    d = etree.Element(qn("a:prstDash"))
    d.set("val", dash)
    fill = ln.find(qn("a:solidFill"))
    if fill is not None:
        fill.addnext(d)
    else:
        ln.insert(0, d)


def add_shape(slide, el):
    r = el["rect"]
    shp = slide.shapes.add_shape(SHAPES[el["shape"]], emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    if el["shape"] == "round" and el["radius_px"]:
        shp.adjustments[0] = min(0.5, el["radius_px"] / min(r["w"], r["h"]))
    if el["fill"]:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(el["fill"]["hex"])
        if el["fill"]["a"] < 1:
            clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(el["fill"]["a"] * 100000)))
    else:
        shp.fill.background()
    if el["line"]:
        shp.line.color.rgb = rgb(el["line"]["hex"])
        shp.line.width = Pt(el["line"]["w"] * 0.75)
        _set_dash(shp.line, el.get("dash"))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    for ref in shp._element.xpath(".//a:effectRef"):  # LibreOffice は effectLst を無視してテーマの影を出す
        ref.set("idx", "0")


def add_line(slide, el):
    r = el["rect"]
    horizontal = r["w"] >= r["h"]
    if el.get("diag"):  # 斜め線: 要素の箱の対角（down＝左上→右下、up＝左下→右上）
        x1, x2 = r["x"], r["x"] + r["w"]
        y1, y2 = (r["y"], r["y"] + r["h"]) if el["diag"] == "down" else (r["y"] + r["h"], r["y"])
        th = el.get("w") or 3
    elif horizontal:
        x1, y1, x2, y2, th = r["x"], r["y"] + r["h"] / 2, r["x"] + r["w"], r["y"] + r["h"] / 2, r["h"]
    else:
        x1, y1, x2, y2, th = r["x"] + r["w"] / 2, r["y"], r["x"] + r["w"] / 2, r["y"] + r["h"], r["w"]
    if el.get("reverse"):
        x1, y1, x2, y2 = x2, y2, x1, y1
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    c.line.color.rgb = rgb(el["fill"]["hex"])
    c.line.width = Pt(max(0.75, th * 0.75))
    _set_dash(c.line, el.get("dash"))
    if el["arrow"]:
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
    for ref in c._element.xpath(".//a:effectRef"):  # テーマの効果（影）を参照しない
        ref.set("idx", "0")


def rasterize_svg(svg: Path, width_px: float, work_name: str) -> Path:
    import pymupdf

    out_dir = paths.temp_dir(work_name) / "raster"
    out_dir.mkdir(parents=True, exist_ok=True)
    # 別のフォルダの同名ファイルと衝突させない
    tag = hashlib.sha1(str(svg.resolve()).encode("utf-8"), usedforsecurity=False).hexdigest()[:8]
    out = out_dir / f"{svg.stem}-{tag}.png"
    doc = pymupdf.open(str(svg))
    page = doc[0]
    zoom = max(1.0, width_px * 3 / page.rect.width)
    page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True).save(str(out))
    doc.close()
    return out


def add_image(slide, el, base_dir, work_name):
    r = el["rect"]
    src = url_to_path(el["src"], base_dir)
    if src.suffix.lower() == ".svg":
        src = rasterize_svg(src, r["w"], work_name)
    pic = slide.shapes.add_picture(str(src), emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    pic._element.nvPicPr.cNvPr.set("descr", el["alt"])


def add_picture(slide, el):
    """絵（pp-picture）: measure が PNG にした画像を、測った位置に貼る。
    名前を pp-picture にして、大きな画像の検査から区別する。"""
    r = el["rect"]
    if not el.get("src"):
        raise ValueError("絵の画像がありません（measure を raster_dir 付きで実行してください）")
    pic = slide.shapes.add_picture(el["src"], emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    pic._element.nvPicPr.cNvPr.set("descr", el["alt"])
    pic._element.nvPicPr.cNvPr.set("name", PICTURE_NAME)


# ---- 表 ----
def _border(tcPr, tag, spec):
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    ln = etree.Element(qn(tag))
    if spec:
        ln.set("w", str(int(spec["w"] * 0.75 * 12700)))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", spec["hex"])
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    return ln


def add_table(slide, el):
    t, r = el["table"], el["rect"]
    xs, rows = t["xs"], t["rows"]
    ncols = len(xs) - 1
    gf = slide.shapes.add_table(len(rows), ncols, emu(r["x"]), emu(r["y"]), emu(r["w"]), emu(r["h"]))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for i in range(ncols):
        tbl.columns[i].width = emu(xs[i + 1] - xs[i])
    for j, row in enumerate(rows):
        tbl.rows[j].height = emu(row["h"])
    for j, row in enumerate(rows):
        for c in row["cells"]:
            cell = tbl.cell(j, c["c0"])
            if c["c1"] - c["c0"] > 1 or c["rowspan"] > 1:
                cell.merge(tbl.cell(j + c["rowspan"] - 1, c["c1"] - 1))
            cell.margin_left, cell.margin_right = emu(c["pad"]["l"]), emu(c["pad"]["r"])
            cell.margin_top, cell.margin_bottom = emu(c["pad"]["t"]), emu(c["pad"]["b"])
            cell.vertical_anchor = VALIGN[c["valign"]]
            if c["fill"]:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(c["fill"]["hex"])
            else:
                cell.fill.background()
            fill_text_frame(cell.text_frame, c["paras"], anchor=VALIGN[c["valign"]])
            tcPr = cell._tc.get_or_add_tcPr()
            for i, (tag, side) in enumerate((("a:lnL", "l"), ("a:lnR", "r"), ("a:lnT", "t"), ("a:lnB", "b"))):
                tcPr.insert(i, _border(tcPr, tag, c["borders"][side]))


# ---- スライド ----
def build(
    model, template_path, info, palette, out: Path, base_dir: Path, footer: str | None = None, work_name: str = "build"
) -> Path:
    from . import charts

    prs = template.open_template(template_path)
    drop_all_slides(prs)
    for s in model:
        layout = prs.slide_layouts[template.layout_index(info, s)]
        slide = prs.slides.add_slide(layout)
        used = set()
        elements = s["elements"]
        # タイトル・副題はレイアウトのプレースホルダーへ
        for el in elements:
            if el["kind"] != "text":
                continue
            is_title = el.get("ph") == "title" or (
                el.get("role") == "title" and s["layout"] not in ("cover", "closing")
            )
            is_sub = el.get("ph") == "subtitle"
            if is_title and slide.shapes.title is not None and "title" not in used:
                ph = slide.shapes.title
                if el.get("ph") is None:  # 本文のタイトルは、測った位置に上書きする
                    _place(ph, el["box"])
                fill_text_frame(ph.text_frame, el["paras"])
                used.add("title")
                el["_done"] = ph
            elif is_sub:
                subs = [p for p in slide.placeholders if p.placeholder_format.type == T.SUBTITLE]
                if subs:
                    fill_text_frame(subs[0].text_frame, el["paras"])
                    used.add("sub")
                    el["_done"] = subs[0]
        for ph in list(slide.placeholders):
            keep = (ph.placeholder_format.type in (T.TITLE, T.CENTER_TITLE) and "title" in used) or (
                ph.placeholder_format.type == T.SUBTITLE and "sub" in used
            )
            if not keep:
                ph._element.getparent().remove(ph._element)
        clone_footer_placeholders(slide, layout, footer)
        for i, el in enumerate(elements):
            ph = el.pop("_done", None)
            if ph is not None:
                if i > 0:  # 前に他の要素がある題・副題は、HTML の重なり順どおり最前面側へ移す
                    ph._element.getparent().append(ph._element)
                continue
            k = el["kind"]
            if k == "text":
                add_textbox(slide, el)
            elif k == "shape":
                add_shape(slide, el)
            elif k == "line":
                add_line(slide, el)
            elif k == "image":
                add_image(slide, el, base_dir, work_name)
            elif k == "picture":
                add_picture(slide, el)
            elif k == "table":
                add_table(slide, el)
            elif k == "chart":
                charts.add_chart(slide, el, palette)
            else:
                raise ValueError(f"未対応の要素: {k}")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
