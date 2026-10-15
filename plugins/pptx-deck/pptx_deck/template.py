"""利用者（または標準）のテンプレートを開いて、使うレイアウトと領域を決める。"""
import io
import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as T
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .pptxutil import PX, SLIDE_H_EMU, SLIDE_W_EMU

_COVER = re.compile(r"(title slide|(?<!裏)表紙|cover|タイトル ?スライド)", re.I)
_CLOSING = re.compile(r"(裏表紙|closing|\bend\b|thank|結び)", re.I)
_BODY = re.compile(r"(title and content|本文|body|タイトルとコンテンツ|content)", re.I)
_SECTION = re.compile(r"(章扉|中扉|セクション|section ?header|chapter|divider|区切り)", re.I)
_AGENDA = re.compile(r"(目次|アジェンダ|agenda|contents?\b|index)", re.I)
# 名前から型への対応を推定するための手がかり（型の名前 → レイアウト名の正規表現）。ここにない型は、カテゴリ（charts・tables・diagrams）の手がかりか、本文で代える
_TYPE_HINTS = {
    "screenshot": re.compile(r"(画像|写真|スクリーン|picture|photo|screenshot)", re.I),
    "org": re.compile(r"(組織|org)", re.I),
    "gantt": re.compile(r"(スケジュール|工程|schedule|timeline|ロードマップ)", re.I),
    "roadmap": re.compile(r"(スケジュール|工程|schedule|timeline|ロードマップ)", re.I),
    "timeline": re.compile(r"(スケジュール|工程|schedule|timeline|ロードマップ)", re.I),
}
_CATEGORY_HINTS = {
    "charts": re.compile(r"(グラフ|chart)", re.I),
    "tables": re.compile(r"(表|table)", re.I),
    "diagrams": re.compile(r"(図解|ダイアグラム|diagram)", re.I),
}
KINDS = ("cover", "section", "agenda", "closing")     # レイアウトの種類（原稿の layout）のうち、テンプレートの専用レイアウトを探すもの


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    w: float
    h: float

    def to_dict(self) -> dict:
        return asdict(self)

    @staticmethod
    def from_dict(d: dict) -> "Rect":
        return Rect(**d)


@dataclass(frozen=True)
class TemplateInfo:
    cover: int
    body: int
    closing: int | None
    area: Rect
    title: Rect
    cover_title: Rect
    cover_subtitle: Rect | None
    closing_title: Rect
    keepout: tuple[Rect, ...] = ()
    notes: tuple[str, ...] = field(default_factory=tuple)
    kinds: dict = field(default_factory=dict)             # 章扉・目次の専用レイアウト {"section": 3, "agenda": 4}。なければ本文で代える
    types: dict = field(default_factory=dict)             # 型ごとの専用レイアウト {"gantt": 7}（data-type が一致するスライドに使う）。なければ本文
    layout_keepout: dict = field(default_factory=dict)    # 専用レイアウトごとの、ヘッダー・フッター・ロゴの領域 {"3": (Rect, …)}

    def to_dict(self) -> dict:
        return {
            "cover": self.cover, "body": self.body, "closing": self.closing,
            "area": self.area.to_dict(), "title": self.title.to_dict(),
            "cover_title": self.cover_title.to_dict(),
            "cover_subtitle": self.cover_subtitle.to_dict() if self.cover_subtitle else None,
            "closing_title": self.closing_title.to_dict(),
            "keepout": [r.to_dict() for r in self.keepout], "notes": list(self.notes),
            "kinds": dict(self.kinds), "types": dict(self.types),
            "layout_keepout": {k: [r.to_dict() for r in v] for k, v in self.layout_keepout.items()},
        }

    @staticmethod
    def from_dict(d: dict) -> "TemplateInfo":
        sub = d["cover_subtitle"]
        return TemplateInfo(
            cover=d["cover"], body=d["body"], closing=d["closing"],
            area=Rect.from_dict(d["area"]), title=Rect.from_dict(d["title"]),
            cover_title=Rect.from_dict(d["cover_title"]),
            cover_subtitle=Rect.from_dict(sub) if sub else None,
            closing_title=Rect.from_dict(d["closing_title"]),
            keepout=tuple(Rect.from_dict(r) for r in d["keepout"]), notes=tuple(d["notes"]),
            kinds=dict(d.get("kinds", {})), types=dict(d.get("types", {})),
            layout_keepout={k: tuple(Rect.from_dict(r) for r in v) for k, v in d.get("layout_keepout", {}).items()})


def layout_index(info: TemplateInfo, slide: dict) -> int:
    """スライド（測ったモデルの 1 枚）に使うレイアウトの番号。型の専用レイアウト → 種類（表紙・章扉・目次・裏表紙）→ 本文の順。"""
    t = slide.get("type") or ""
    if t in info.types:
        return info.types[t]
    kind = slide.get("layout")
    if kind == "cover":
        return info.cover
    if kind == "closing":
        return info.closing if info.closing is not None else info.body
    if kind in info.kinds:
        return info.kinds[kind]
    return info.body


def keepout_for(info: TemplateInfo, index: int) -> tuple[Rect, ...]:
    """そのレイアウトの、本文を置いてはいけない領域。専用レイアウトで測っていなければ、本文のもの。"""
    return info.layout_keepout.get(str(index), info.keepout)


def open_template(path) -> Presentation:
    """.pptx / .potx を開く。.potx は content-type を .pptx 相当に直して開く。"""
    path = Path(path)
    if path.suffix.lower() == ".potx":
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                zout.writestr(it, data)
        buf.seek(0)
        prs = Presentation(buf)
    else:
        prs = Presentation(str(path))
    if abs(prs.slide_width - SLIDE_W_EMU) > SLIDE_W_EMU * 0.01 or \
            abs(prs.slide_height - SLIDE_H_EMU) > SLIDE_H_EMU * 0.01:
        raise ValueError(
            "テンプレートは 13.33×7.5in（16:9）のみ対応です。"
            f"実際: {prs.slide_width / 914400:.2f}×{prs.slide_height / 914400:.2f}in")
    return prs


def _rect(shape) -> Rect:
    return Rect(shape.left / PX, shape.top / PX, shape.width / PX, shape.height / PX)


def _phs(layout, *types):
    return [p for p in layout.placeholders if p.placeholder_format.type in types]


def _pick(prs, name_re, pred=None):
    for i, lay in enumerate(prs.slide_layouts):
        if name_re.search(lay.name):
            return i
    if pred:
        for i, lay in enumerate(prs.slide_layouts):
            if pred(lay):
                return i
    return None


def _fixed_shapes(layout) -> tuple[Rect, ...]:
    """本文レイアウトとマスターにある、プレースホルダーでない図形（ロゴ・帯など）。全面の背景と面積 0 の線は除く。"""
    out = []
    for shapes in (layout.shapes, layout.slide_master.shapes):
        for sh in shapes:
            if sh.is_placeholder or None in (sh.left, sh.top, sh.width, sh.height):
                continue
            r = _rect(sh)
            if r.w * r.h <= 0 or r.w * r.h > 1280 * 720 * 0.5:
                continue
            out.append(r)
    return tuple(out)


def analyze(path) -> TemplateInfo:
    prs = open_template(path)
    L = list(prs.slide_layouts)
    notes = []
    body = _pick(prs, _BODY, lambda l: _phs(l, T.TITLE) and _phs(l, T.BODY, T.OBJECT))
    if body is None:
        raise ValueError("本文に使えるレイアウト（タイトルと本文の領域があるもの）がありません")
    cover = _pick(prs, _COVER, lambda l: bool(_phs(l, T.CENTER_TITLE)))
    if cover is None:
        cover = body
        notes.append("表紙のレイアウトが見つかりません。本文のレイアウトで代えます。")
    closing = _pick(prs, _CLOSING)
    if closing is None:
        notes.append("裏表紙のレイアウトが見つかりません。本文のレイアウトで代えます。")
    b, c = L[body], L[cover]
    cl_titles = _phs(L[closing], T.TITLE, T.CENTER_TITLE) if closing is not None else []
    areas = _phs(b, T.BODY, T.OBJECT)
    area = _rect(areas[0]) if areas else Rect(64, 150, 1152, 490)
    title = _rect(_phs(b, T.TITLE)[0])
    ct = _phs(c, T.CENTER_TITLE, T.TITLE)
    cs = _phs(c, T.SUBTITLE)
    keepout = tuple(_rect(p) for p in _phs(b, T.SLIDE_NUMBER, T.FOOTER, T.DATE)) + _fixed_shapes(b)
    kinds, types, lay_keep, extra_notes = _estimate_layouts(prs, cover, body, closing)
    return TemplateInfo(
        cover=cover, body=body, closing=closing, area=area, title=title,
        cover_title=_rect(ct[0]) if ct else title,
        cover_subtitle=_rect(cs[0]) if cs else None,
        closing_title=_rect(cl_titles[0]) if cl_titles else title,
        keepout=keepout, notes=tuple(notes + extra_notes), kinds=kinds, types=types, layout_keepout=lay_keep)


def component_types() -> dict[str, str]:
    """型の名前 → カテゴリ（assets/components の下のフォルダ名）。デザイン 2・3（-d2・-d3）は同じ型。"""
    from . import paths
    out = {}
    root = paths.ASSETS / "components"
    if root.is_dir():
        for f in sorted(root.rglob("*.html")):
            out.setdefault(re.sub(r"-d[23]$", "", f.stem), f.parent.name)
    return out


def _usable(layout) -> bool:
    return bool(_phs(layout, T.TITLE, T.CENTER_TITLE))


def _layout_keepout(layout) -> tuple[Rect, ...]:
    return tuple(_rect(p) for p in _phs(layout, T.SLIDE_NUMBER, T.FOOTER, T.DATE)) + _fixed_shapes(layout)


def _estimate_layouts(prs, cover, body, closing):
    """テンプレートの全レイアウトから、章扉・目次・型ごとの専用レイアウトを名前で推定する。
    推定できない種類・型は入れない（本文のレイアウトで代える）。戻り値は（kinds, types, 専用レイアウトの keepout, notes）。"""
    L = list(prs.slide_layouts)
    taken = {cover, body, closing}
    kinds: dict[str, int] = {}
    for kind, rx in (("section", _SECTION), ("agenda", _AGENDA)):
        for i, lay in enumerate(L):
            if i not in taken and _usable(lay) and rx.search(lay.name):
                kinds[kind] = i
                taken.add(i)
                break
    types: dict[str, int] = {}
    for t, cat in component_types().items():
        if t in KINDS:
            continue
        rx = _TYPE_HINTS.get(t) or _CATEGORY_HINTS.get(cat)
        if not rx:
            continue
        for i, lay in enumerate(L):
            if i not in (cover, body, closing) and i not in kinds.values() and _usable(lay) and rx.search(lay.name):
                types[t] = i
                break
    notes = []
    for kind, label in (("section", "章扉"), ("agenda", "目次")):
        if kind not in kinds:
            notes.append(f"{label}のレイアウトが見つかりません。本文のレイアウトで代えます。")
    used = set(kinds.values()) | set(types.values())
    return kinds, types, {str(i): _layout_keepout(L[i]) for i in sorted(used)}, notes


def describe(path, info: TemplateInfo) -> dict:
    """利用者に確認してもらう対応表。レイアウトの一覧（名前とプレースホルダー）と、型ごとの対応（本文で代える型は 1 行にまとめる）。"""
    prs = open_template(path)
    L = list(prs.slide_layouts)
    name = lambda i: L[i].name
    layouts = [{"index": i, "name": lay.name,
                "placeholders": [str(p.placeholder_format.type).split(".")[-1].split(" ")[0] for p in lay.placeholders]}
               for i, lay in enumerate(L)]
    rows = [{"key": "cover", "layout": name(info.cover), "index": info.cover},
            {"key": "section", "layout": name(info.kinds["section"]) if "section" in info.kinds else None, "index": info.kinds.get("section")},
            {"key": "agenda", "layout": name(info.kinds["agenda"]) if "agenda" in info.kinds else None, "index": info.kinds.get("agenda")},
            {"key": "closing", "layout": name(info.closing) if info.closing is not None else None, "index": info.closing}]
    for t in sorted(info.types):
        rows.append({"key": t, "layout": name(info.types[t]), "index": info.types[t]})
    all_types = [t for t in component_types() if t not in KINDS]
    return {"body": {"layout": name(info.body), "index": info.body}, "layouts": layouts, "rows": rows,
            "body_types": [t for t in all_types if t not in info.types]}


def set_mapping(path, info: TemplateInfo, key: str, layout_name: str | None) -> TemplateInfo:
    """利用者の修正を反映する。key は cover・section・agenda・closing か型の名前。layout_name が None なら専用レイアウトをやめて本文で代える。
    layout_name はレイアウトの名前（完全一致を優先、なければ部分一致が 1 つだけのとき）か番号。"""
    from dataclasses import replace
    prs = open_template(path)
    L = list(prs.slide_layouts)
    valid = set(KINDS) | set(component_types())
    if key not in valid:
        raise ValueError(f"{key} は対応づけできません（cover・section・agenda・closing か、型の名前: components.md）")
    idx = None
    if layout_name is not None:
        if layout_name.isdigit() and int(layout_name) < len(L):
            idx = int(layout_name)
        else:
            exact = [i for i, lay in enumerate(L) if lay.name == layout_name]
            part = [i for i, lay in enumerate(L) if layout_name in lay.name]
            hits = exact or part
            if len(hits) != 1:
                raise ValueError(f"レイアウト「{layout_name}」が決まりません（候補: {', '.join(lay.name for lay in L)}）")
            idx = hits[0]
        if not _usable(L[idx]):
            raise ValueError(f"レイアウト「{L[idx].name}」にタイトルの領域がないため使えません")
    kinds, types = dict(info.kinds), dict(info.types)
    cover, closing = info.cover, info.closing
    if key == "cover":
        cover = idx if idx is not None else info.body
    elif key == "closing":
        closing = idx
    elif key in ("section", "agenda"):
        kinds.pop(key, None)
        if idx is not None:
            kinds[key] = idx
    else:
        types.pop(key, None)
        if idx is not None:
            types[key] = idx
    used = set(kinds.values()) | set(types.values())
    lay_keep = {str(i): _layout_keepout(L[i]) for i in sorted(used)}
    c = L[cover]
    ct = _phs(c, T.CENTER_TITLE, T.TITLE)
    cs = _phs(c, T.SUBTITLE)
    cl = _phs(L[closing], T.TITLE, T.CENTER_TITLE) if closing is not None else []
    return replace(info, cover=cover, closing=closing, kinds=kinds, types=types, layout_keepout=lay_keep,
                   cover_title=_rect(ct[0]) if ct else info.title,
                   cover_subtitle=_rect(cs[0]) if cs else None,
                   closing_title=_rect(cl[0]) if cl else info.title)


_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
       "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
_THEME_KEYS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")


def _theme_colors(master) -> dict[str, str]:
    """テーマの色（dk1・lt1・accent1…）→ 6 桁の 16 進。"""
    try:
        root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    except KeyError:                                     # テーマ part がない。srgbClr の直接指定だけ通す
        return {}
    out = {}
    for key in _THEME_KEYS:
        el = root.find(f".//a:clrScheme/a:{key}", _NS)
        if el is None or len(el) == 0:
            continue
        c = el[0]
        val = c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr")
        if val:
            out[key] = val.upper()
    return out


def _color_hex(el, theme: dict[str, str], clrmap: dict[str, str]) -> str | None:
    if len(el):                                          # lumMod・tint・alpha などの変換子付きは正しい色が決められない
        return None
    if el.tag.endswith("srgbClr"):
        return el.get("val").upper()
    if el.tag.endswith("schemeClr"):
        key = el.get("val")
        return theme.get(clrmap.get(key, key))
    if el.tag.endswith("sysClr"):
        return (el.get("lastClr") or "FFFFFF").upper()
    return None


def layout_bg_hex(prs, layout_index: int) -> str | None:
    """レイアウト（なければマスター）の背景が単色なら '#RRGGBB'。画像・グラデーション・取れないときは None。"""
    layout = prs.slide_layouts[layout_index]
    master = layout.slide_master
    theme = _theme_colors(master)
    cm = master._element.find("p:clrMap", _NS)
    clrmap = dict(cm.attrib) if cm is not None else {}
    ovr = layout._element.find("p:clrMapOvr/a:overrideClrMapping", _NS)
    if ovr is not None:                                  # レイアウトの色の対応づけがマスターより優先（masterClrMapping ならマスターのまま）
        clrmap = dict(ovr.attrib)
    for part in (layout._element, master._element):
        bg = part.find("p:cSld/p:bg", _NS)
        if bg is None:
            continue
        solid = bg.find("p:bgPr/a:solidFill", _NS)
        if solid is not None and len(solid):
            h = _color_hex(solid[0], theme, clrmap)
            return f"#{h}" if h else None
        ref = bg.find("p:bgRef", _NS)
        if ref is not None:                              # テーマの背景スタイルの参照
            if ref.get("idx") != "1001" or not len(ref):  # 1001 はテーマの 1 番目の塗り（通常は単色）。1002 以上は単色と言い切れない
                return None
            h = _color_hex(ref[0], theme, clrmap)
            return f"#{h}" if h else None
        return None                                      # 画像・グラデーションなど
    return None


def background_colors(path, info: TemplateInfo) -> tuple[dict[str, str], list[str]]:
    """表紙・本文・裏表紙の背景色（HTML の下地）。単色でないレイアウトは白で近似し、notes に理由を返す。"""
    prs = open_template(path)
    idx = {"cover": info.cover, "body": info.body, "closing": info.closing if info.closing is not None else info.body}
    idx.update(info.kinds)
    idx.update({f"type:{t}": i for t, i in info.types.items()})
    out, notes = {}, []
    for name, i in idx.items():
        h = layout_bg_hex(prs, i)
        if h is None:
            h = "#FFFFFF"
            notes.append(f"{name} の背景が単色ではないため、HTML の下地は白で近似しました（PPTX にはテンプレートの背景が出ます）。")
        out[name] = h
    return out, notes


def layout_css(info: TemplateInfo, bg_hex: str, layout_bg: dict[str, str] | None = None) -> str:
    def v(prefix, r: Rect):
        return (f"  --{prefix}-x: {r.x:g}px;\n  --{prefix}-y: {r.y:g}px;\n"
                f"  --{prefix}-w: {r.w:g}px;\n  --{prefix}-h: {r.h:g}px;\n")
    root = f":root {{\n  --bg: {bg_hex};\n" + v("area", info.area) + v("title", info.title) \
        + v("cover-title", info.cover_title) + v("closing-title", info.closing_title)
    if info.cover_subtitle:
        root += v("cover-subtitle", info.cover_subtitle)
    root += "}\n"
    for name in ("cover", "section", "agenda", "closing"):       # 表紙・章扉・目次・裏表紙の下地が本文と違う色でも、文字色の検査が合うようにする
        if layout_bg and layout_bg.get(name) and layout_bg[name] != bg_hex:
            root += f'.slide[data-layout="{name}"] {{ --bg: {layout_bg[name]}; }}\n'
    for name, color in (layout_bg or {}).items():           # 型ごとの専用レイアウトの下地（data-type が一致するスライド）
        if name.startswith("type:") and color != bg_hex:
            root += f'.slide[data-type="{name[5:]}"] {{ --bg: {color}; }}\n'
    return root
