"""体裁の検査（最小限）: はみ出し・重なり・文字サイズの下限・コントラスト・画像の使い方。"""

import re
from pathlib import Path

from . import template, theme
from .pptxutil import file_url_to_path
from .template import TemplateInfo

SLIDE_W, SLIDE_H = 1280, 720
MIN_PT_DEFAULT, MIN_PT_SMALL = 14, 10  # 本文 / 注記（role=caption|chrome）
LARGE_PX = 24  # 24px(18pt) 以上は大きい文字（3:1）
BODY_IMAGE_RATIO = 0.4
OVERLAP_MIN_AREA = 4
AREA_BOTTOM_GAP = 120  # 本文の下端が領域下端よりこれ以上上なら空白帯
SAME_WORD_MIN = 6  # 同じ文言の重複を数える最小字数
_DASH = re.compile("[—―─]")
SHAPE_KINDS = {"rect", "round", "ellipse", "triangle"}  # build.py の SHAPES と同じ。これ以外は build が止まる
PICTURE_NAME = "pp-picture"  # build.py の PICTURE_NAME と同じ（絵として貼った画像の名前）
TITLE_MAX, LEAD_MIN, LEAD_MAX = (
    20,
    30,
    60,
)  # 題（短い名詞句）の字数の上限、ヘッドメッセージ（判断を言い切る 1 文）の字数


def _v(rule, slide, message, severity="error"):
    return {"rule": rule, "slide": slide, "message": message, "severity": severity}


def errors(violations):
    return [v for v in violations if v["severity"] == "error"]


def _r(el):
    return el.get("box") if el["kind"] == "text" else el["rect"]


def _inter(a, b):
    ix = min(a["x"] + a["w"], b["x"] + b["w"]) - max(a["x"], b["x"])
    iy = min(a["y"] + a["h"], b["y"] + b["h"]) - max(a["y"], b["y"])
    return ix * iy if ix > 0 and iy > 0 else 0


def _contains(a, b, tol=2):
    return (
        a["x"] - tol <= b["x"]
        and a["y"] - tol <= b["y"]
        and a["x"] + a["w"] + tol >= b["x"] + b["w"]
        and a["y"] + a["h"] + tol >= b["y"] + b["h"]
    )


def _check_picture(e, n, label):
    """絵（pp-picture）の検査。絵の中身は PPTX で直せないので、HTML の段階で文字・色・代替テキストを確かめる。
    絵の内部の重なり（矢印・番号・吹き出しを絵に重ねる）は検査しない。絵全体とほかの部品の重なりは、呼び出し側が検査する。"""
    out = []
    if not e.get("alt"):
        out.append(
            _v("imageAlt", n, f"{label}: 絵の代替テキスト（data-alt）がありません。絵の内容を 1〜2 文で書いてください")
        )
    if not e.get("hasSvg"):
        out.append(_v("pictureSvg", n, f"{label}: pp-picture の中に <svg> がありません"))
    if e.get("badColors"):
        out.append(
            _v(
                "pictureColor",
                n,
                f"{label}: テーマの変数以外の色があります（{', '.join(e['badColors'][:4])}）。"
                "var(--accent) などの変数か、sv-f-* / sv-s-* のクラスを使ってください",
                "warning",
            )
        )
    texts = e.get("texts", [])
    for t in texts:
        floor = MIN_PT_SMALL if t.get("small") else MIN_PT_DEFAULT
        pt = t["size_px"] * 0.75
        if pt < floor:
            out.append(
                _v(
                    "fontMin",
                    n,
                    f"{label}: 絵の中の文字「{t['text'][:10]}」は {pt:.1f}pt で、下限 {floor}pt 未満です"
                    "（図を 2 枚に分けるか、番号だけを絵に置いて説明を凡例の表に出す）",
                )
            )
        if t.get("outside"):
            out.append(_v("pictureOverflow", n, f"{label}: 絵の中の文字「{t['text'][:10]}」が絵の外にはみ出しています"))
    for i, a in enumerate(texts):
        for b in texts[i + 1 :]:
            if _inter(a["rect"], b["rect"]) > OVERLAP_MIN_AREA:
                out.append(
                    _v(
                        "pictureTextOverlap",
                        n,
                        f"{label}: 絵の中の文字「{a['text'][:8]}」と「{b['text'][:8]}」が重なっています",
                        "warning",
                    )
                )
    return out


def check_model(model, info: TemplateInfo):
    out = []
    body_keepout = [r.to_dict() for r in info.keepout]
    for s in model:
        n = s["index"]
        # 専用レイアウト（章扉・目次・型ごと）を使うスライドは、そのレイアウトのヘッダー・フッター・ロゴの領域で検査する
        li = template.layout_index(info, s)
        keepout = (
            [r.to_dict() for r in template.keepout_for(info, li)]
            if li in (*info.kinds.values(), *info.types.values())
            else body_keepout
        )
        els = s["elements"]
        has_title = any(e["kind"] == "text" and (e.get("role") == "title" or e.get("ph") == "title") for e in els)
        if not has_title and s["layout"] != "closing":
            out.append(_v("titleMissing", n, 'タイトル（data-role="title" または data-ph="title"）がありません'))
        for e in els:
            if e["kind"] == "shape" and e.get("shape") not in SHAPE_KINDS:
                out.append(
                    _v(
                        "shapeKind",
                        n,
                        f'shape#{e["idx"]}: data-shape="{e.get("shape")}" は使えません'
                        f"（{' / '.join(sorted(SHAPE_KINDS))}）",
                    )
                )
        if s["layout"] in ("content", "diagram", "table", "chart"):
            leads = [e for e in els if e["kind"] == "text" and e.get("role") == "lead"]
            if not leads:
                out.append(
                    _v(
                        "leadMissing",
                        n,
                        'ヘッドメッセージ（data-role="lead"。判断を言い切る 1 文、30〜60 字）がありません',
                        "warning",
                    )
                )
            else:
                lt = "".join(r.get("text", "") for pa in leads[0]["paras"] for r in pa["runs"]).strip()
                if not LEAD_MIN <= len(lt) <= LEAD_MAX:
                    out.append(
                        _v(
                            "leadLength",
                            n,
                            f"ヘッドメッセージが {len(lt)} 字です（{LEAD_MIN}〜{LEAD_MAX} 字。2 行までに収める）",
                            "warning",
                        )
                    )
            ttl = [e for e in els if e["kind"] == "text" and (e.get("role") == "title" or e.get("ph") == "title")]
            if ttl:
                tt = "".join(r.get("text", "") for pa in ttl[0]["paras"] for r in pa["runs"]).strip()
                if len(tt) > TITLE_MAX:
                    out.append(
                        _v(
                            "titleLength",
                            n,
                            f"題が {len(tt)} 字です（{TITLE_MAX} 字以内の短い名詞句。判断はヘッドメッセージに書く）",
                            "warning",
                        )
                    )
        texts = []
        for e in els:
            if e["kind"] == "text":
                t = "".join(r.get("text", "") for p in e["paras"] for r in p["runs"]).strip()
                texts.append((e, t))
                if _DASH.search(t):
                    out.append(_v("dash", n, f"text#{e['idx']}: ダッシュ（—）を使わないでください", "warning"))
                if e.get("lastLineChars") is not None and e["lastLineChars"] <= 2:
                    out.append(
                        _v(
                            "orphanLine",
                            n,
                            f"text#{e['idx']}: 最終行が {e['lastLineChars']} 字だけです。幅か文を直してください",
                            "warning",
                        )
                    )
        seen = {}
        for e, t in texts:
            if len(t) >= SAME_WORD_MIN and not e.get("chrome"):
                seen.setdefault(t, []).append(e["idx"])
        for t, idxs in seen.items():
            if len(idxs) > 1:
                out.append(
                    _v(
                        "sameWord",
                        n,
                        f"同じ文言「{t[:12]}」が 1 枚に {len(idxs)} 回あります。軸・見出しにくくり出してください",
                        "warning",
                    )
                )
        if s["layout"] in ("content", "diagram", "table", "chart"):
            body = [
                e
                for e in els
                if not e.get("chrome") and not e.get("ph") and e.get("role") != "title" and e["kind"] != "line"
            ]
            if body:
                bottom = max(_r(e)["y"] + _r(e)["h"] for e in body)
                area_bottom = max((z["y"] for z in keepout), default=SLIDE_H)
                if bottom < area_bottom - AREA_BOTTOM_GAP:
                    out.append(
                        _v(
                            "emptyBand",
                            n,
                            f"本文が下端（y {area_bottom}）の手前 {area_bottom - bottom:.0f}px で終わっています。"
                            "無所属の空白帯を残さない",
                            "warning",
                        )
                    )
        for e in els:
            k, rect = e["kind"], _r(e)
            label = f"{k}#{e['idx']}"
            if k == "text":
                if e.get("overflow"):
                    out.append(
                        _v(
                            "overflow",
                            n,
                            f"{label}: 文字が枠からはみ出しています。幅・高さを広げるか文章を減らしてください",
                        )
                    )
                if e.get("lost"):
                    out.append(
                        _v(
                            "textLost",
                            n,
                            f"{label}: 「{e['lost'][:20]}」が PPTX に入りません。"
                            "段落は p・div・li で並べ、入れ子の ul や段落と混ざった裸の文字を使わないでください",
                        )
                    )
                for p in e["paras"]:
                    for run in p["runs"]:
                        if "text" not in run:
                            continue
                        pt = run["size_px"] * 0.75
                        small = e.get("role") in ("caption", "chrome") or e.get("chrome")
                        floor = MIN_PT_SMALL if small else MIN_PT_DEFAULT
                        if pt < floor:
                            out.append(_v("fontMin", n, f"{label}: {pt:g}pt は下限 {floor}pt 未満です"))
                        need = 3.0 if run["size_px"] >= LARGE_PX else 4.5
                        ratio = theme.contrast("#" + run["color"], "#" + e["bg"])
                        if ratio < need:
                            out.append(
                                _v(
                                    "contrast",
                                    n,
                                    f"{label}: 文字色 #{run['color']} と背景 #{e['bg']} の"
                                    f"コントラスト比 {ratio:.1f} が {need} 未満です",
                                )
                            )
            if k == "chart":
                from . import charts

                for msg in charts.validate(e["chart"]):
                    out.append(_v("chartSpec", n, f"{label}: {msg}"))
            if k == "image":
                src = e.get("src", "")
                if not src.startswith("file:"):
                    out.append(
                        _v(
                            "imageSrc",
                            n,
                            f"{label}: 画像の src がローカルのファイルではありません（{src[:60]}）。"
                            "相対パスのファイルを指定してください",
                        )
                    )
                elif not file_url_to_path(src).exists():
                    out.append(_v("imageSrc", n, f"{label}: 画像のファイルが見つかりません（{src[:60]}）"))
            if k == "picture":
                out.extend(_check_picture(e, n, label))
            if k == "table":
                for row in e["table"]["rows"]:
                    for c in row["cells"]:
                        for p in c["paras"]:
                            for run in p["runs"]:
                                if "text" in run and run["size_px"] * 0.75 < MIN_PT_DEFAULT:
                                    out.append(
                                        _v(
                                            "fontMin",
                                            n,
                                            f"{label}: 表の文字 {run['size_px'] * 0.75:g}pt は"
                                            f"下限 {MIN_PT_DEFAULT}pt 未満です",
                                        )
                                    )
                                bg = c["fill"]["hex"] if c["fill"] else c["bg"]
                                need = 3.0 if run.get("size_px", 0) >= LARGE_PX else 4.5
                                if "text" in run and theme.contrast("#" + run["color"], "#" + bg) < need:
                                    out.append(
                                        _v(
                                            "contrast",
                                            n,
                                            f"{label}: 表の文字色 #{run['color']} と背景 #{bg} の"
                                            f"コントラスト比が {need} 未満です",
                                        )
                                    )
            if k == "image":
                if not e.get("alt"):
                    out.append(_v("imageAlt", n, f"{label}: alt がありません"))
                area = rect["w"] * rect["h"]
                if e.get("role") not in ("icon", "logo", "photo") or area > SLIDE_W * SLIDE_H * BODY_IMAGE_RATIO:
                    out.append(
                        _v(
                            "bodyImage",
                            n,
                            f"{label}: 本文を画像にしないでください（画像はアイコン・ロゴ・利用者の写真だけ）",
                        )
                    )
            if (
                rect["x"] < -1
                or rect["y"] < -1
                or rect["x"] + rect["w"] > SLIDE_W + 1
                or rect["y"] + rect["h"] > SLIDE_H + 1
            ):
                out.append(_v("outOfSlide", n, f"{label}: スライドの外にはみ出しています"))
            if not e.get("chrome") and not e.get("ph") and k != "line":
                for z in keepout:
                    if _inter(rect, z) > OVERLAP_MIN_AREA:
                        out.append(_v("keepout", n, f"{label}: テンプレートのヘッダー・フッター領域に入っています"))
                        break
        # 重なり
        cand = [
            e
            for e in els
            if e["kind"] in ("text", "shape", "image", "table", "chart", "picture")
            and not (e["kind"] == "text" and not e["paras"])
        ]
        for i, a in enumerate(cand):
            for b in cand[i + 1 :]:
                if a["overlap"] or b["overlap"]:
                    continue
                ra, rb = _r(a), _r(b)
                if _inter(ra, rb) <= OVERLAP_MIN_AREA:
                    continue
                if (a["kind"] == "shape" and _contains(ra, rb)) or (b["kind"] == "shape" and _contains(rb, ra)):
                    continue
                out.append(
                    _v(
                        "overlap",
                        n,
                        f"{a['kind']}#{a['idx']} と {b['kind']}#{b['idx']} が重なっています"
                        "（意図的なら data-overlap を両方に付ける）",
                    )
                )
    return out


def check_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    tol = 9525
    out = []
    for i, s in enumerate(prs.slides, start=1):
        for sh in s.shapes:
            if sh.left is None:
                continue
            if sh.left < -tol or sh.top < -tol or sh.left + sh.width > W + tol or sh.top + sh.height > H + tol:
                out.append(_v("pptxBounds", i, f"{sh.name}: スライドの外にはみ出しています"))
            if sh.shape_type == 13 and sh.name != PICTURE_NAME and sh.width * sh.height > W * H * BODY_IMAGE_RATIO:
                out.append(_v("pptxPicture", i, f"{sh.name}: 大きな画像があります（本文を画像にしない）"))
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.font.size.pt < MIN_PT_SMALL:
                            out.append(
                                _v(
                                    "pptxFontMin",
                                    i,
                                    f"{sh.name}: {r.font.size.pt:g}pt は下限 {MIN_PT_SMALL}pt 未満です",
                                )
                            )
    return out
