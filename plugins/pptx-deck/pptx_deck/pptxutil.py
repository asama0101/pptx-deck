import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525
FONT = "BIZ UDPゴシック"
SLIDE_W_EMU = 12192000  # 13.333in
SLIDE_H_EMU = 6858000   # 7.5in


def emu(px: float) -> Emu:
    return Emu(round(px * PX))


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.lstrip("#").upper())


def drop_all_slides(prs) -> None:
    """テンプレートの既存スライドをすべて外す（マスターとレイアウトは残す）。"""
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)
    # セクション・目的別スライド ショーが、消したスライドの id を指したままだと PowerPoint が修復を求める
    root = prs.part._element
    for ref in list(root.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId")):
        ref.getparent().remove(ref)
    for show in root.findall(qn("p:custShowLst")):
        root.remove(show)


def _next_id(slide) -> int:
    ids = [int(x) for x in slide.shapes._spTree.xpath("//p:cNvPr/@id")]
    return max(ids + [1]) + 1


def clone_footer_placeholders(slide, layout, footer=None) -> None:
    """レイアウトのフッター（コピーライト）・スライド番号のプレースホルダーをスライドに写す（python-pptx は写さない）。
    日付は出さない（ページ枠は、左下にコピーライト、右下にページ番号）。footer を渡すと、フッターの文言をそれに差し替える。"""
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    for ph in layout.placeholders:
        if ph.placeholder_format.type not in (T.SLIDE_NUMBER, T.FOOTER):
            continue
        el = copy.deepcopy(ph._element)
        el.xpath(".//p:cNvPr")[0].set("id", str(_next_id(slide)))
        slide.shapes._spTree.append(el)
        if footer is not None and ph.placeholder_format.type == T.FOOTER:
            ts = el.xpath(".//a:t")
            if ts:
                ts[0].text = footer
                for extra in ts[1:]:
                    extra.text = ""
            else:
                p = el.xpath(".//a:p")[0]
                r = etree.Element(qn("a:r"))
                etree.SubElement(r, qn("a:t")).text = footer
                end = p.find(qn("a:endParaRPr"))
                if end is not None:      # endParaRPr は段落の最後に置く
                    end.addprevious(r)
                else:
                    p.append(r)


def set_run_font(run, size_px: float, bold: bool, italic: bool, color_hex: str) -> None:
    f = run.font
    f.size = Pt(round(size_px * 0.75 * 2) / 2)  # 0.5pt 刻み
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color_hex)
    f.name = FONT  # a:latin
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree
        ea = etree.Element(qn("a:ea"))
        latin.addnext(ea)
    ea.set("typeface", FONT)
