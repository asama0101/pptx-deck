"""題・副題のプレースホルダーは、HTML の要素の順（重なり順）に置く。"""

# ruff: noqa: S101  # pytest の assert
import pytest
from pptx import Presentation
from pptx_deck import build, template

from tests._measure import model_of

pytestmark = pytest.mark.slow  # ブラウザ（Chromium）で測る

BAND = (
    '<div class="pp-shape" data-role="band" style="position:absolute;left:0;top:0;'
    'width:1280px;height:200px;background:#123456"></div>'
)
TITLE = (
    '<div class="pp-text" data-role="title" style="position:absolute;left:40px;top:40px;'
    'width:800px;height:60px;color:#fff;font-size:32px">題</div>'
)


def _texts_in_z_order(inner, tpl, tmp_path):
    """HTML 断片を PPTX にして、1 枚目の図形の文字を背面から前面の順に返す（文字のない図形は None）。"""
    info = template.analyze(tpl)
    out = build.build(model_of(inner), tpl, info, None, tmp_path / "o.pptx", base_dir=tmp_path)
    prs = Presentation(str(out))
    return [s.text_frame.text if s.has_text_frame else None for s in prs.slides[0].shapes]


@pytest.mark.parametrize(
    ("inner", "title_in_front"),
    [
        (BAND + TITLE, True),
        (TITLE + BAND, False),
        (TITLE, False),
    ],
    ids=["帯のあとの題は前面", "帯の前の題は背面", "題だけなら先頭"],
)
def test_題は_HTML_の要素の順に重なる(pkg, tmp_path, inner, title_in_front):
    tpl = pkg / "assets" / "template" / "standard.potx"
    texts = _texts_in_z_order(inner, tpl, tmp_path)
    assert texts.count("題") == 1
    assert texts.index("題") == (len(texts) - 1 if title_in_front else 0)
