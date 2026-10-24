"""縦の破線・点線は border-left で描き、測定でも PPTX でも破線になる。"""

# ruff: noqa: S101  # pytest の assert
import pytest
from lxml import etree
from pptx import Presentation
from pptx_deck import build, template

from tests._measure import model_of

pytestmark = pytest.mark.slow  # ブラウザ（Chromium）で測る

VARS = ":root{--accent:#1F4E79;--text-sub:#666666;--line:#CCCCCC}"


@pytest.fixture
def tpl(pkg):
    return pkg / "assets" / "template" / "standard.potx"


def line_elements(inner):
    slide = model_of(inner, VARS)[0]
    return [e for e in slide["elements"] if e["kind"] == "line"]


@pytest.mark.slow
def test_縦の破線は幅2の線として測られる():
    (e,) = line_elements('<div class="pp-line dashed vl" style="left:100px;top:100px;height:200px"></div>')
    assert e["dash"] == "dash"
    assert e["w"] == 2
    assert e["rect"]["h"] == 200
    assert e["rect"]["w"] == 2  # width:0 + border-left 2px。border-top 描きだと 0 になる
    assert e["fill"]["hex"] == "1F4E79"


@pytest.mark.slow
def test_縦の点線はsysDotとして測られる():
    (e,) = line_elements('<div class="pp-line dotted vl" style="left:100px;top:100px;height:200px"></div>')
    assert e["dash"] == "sysDot"


@pytest.mark.slow
def test_横の破線は従来どおり():
    (e,) = line_elements('<div class="pp-line dashed" style="left:100px;top:100px;width:200px"></div>')
    assert e["dash"] == "dash"
    assert e["w"] == 2


@pytest.mark.slow
@pytest.mark.parametrize(("cls", "prst"), [("dashed", "dash"), ("dotted", "sysDot")], ids=["破線", "点線"])
def test_縦線のPPTXに破線の指定が出る(tmp_path, tpl, cls, prst):
    model = model_of(f'<div class="pp-line {cls} vl" style="left:100px;top:100px;height:200px"></div>', VARS)
    out = build.build(model, tpl, template.analyze(tpl), None, tmp_path / "o.pptx", base_dir=tmp_path)
    xml = "".join(etree.tostring(s._element, encoding="unicode") for s in Presentation(str(out)).slides[0].shapes)
    assert f'prstDash val="{prst}"' in xml
