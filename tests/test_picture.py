"""絵（pp-picture）: SVG を Chromium で PNG にして PPTX に貼る・検査する。

build・check は Chromium を実際に起動して測る（slow）。
"""

# ruff: noqa: S101  # pytest の assert
import io
import json
import re
import xml.etree.ElementTree as ET

import pytest
from pptx_deck import paths, svgids

SAMPLE_NAMES = ["traffic-flow", "network", "failure", "self-intro", "annotate"]

slow = pytest.mark.slow


@pytest.fixture
def run(capsys):
    """cli.main を実行して (終了コード, 標準出力, 標準エラー出力) を返す関数。"""
    from pptx_deck import cli

    def _run(*argv):
        capsys.readouterr()
        rc = cli.main(list(argv))
        captured = capsys.readouterr()
        return rc, captured.out, captured.err

    return _run


@pytest.fixture
def pictures(pkg):
    return pkg / "assets" / "components" / "pictures"


@pytest.fixture
def sample(pictures):
    return (pictures / "traffic-flow.html").read_text(encoding="utf-8")


@pytest.fixture
def make_work(tmp_path, run):
    """init → 断片を 01-A に置く → finalize。final.html のある作業フォルダを返す関数（base は tmp_path/demo）。"""
    base = tmp_path / "out" / "pptx-deck-source" / "demo"

    def _make(fragment):
        assert run("init", str(base), "--color", "navy")[0] == 0
        work = paths.resolve_work(base)
        (work / "slides" / "01-A.html").write_text(fragment, encoding="utf-8")
        (work / "answers-1.md").write_text(
            "---\n---\n## S01\n- 選択: A\n- 判定: 承認\n- コメント: \n", encoding="utf-8"
        )
        assert run("finalize", str(base))[0] == 0
        return work

    _make.base = base
    return _make


def init_with_samples(run, base, pictures, names, answer_tail="- 判定: 承認\n- コメント: \n"):
    """init して names の見本を 01-A, 02-A… に置き、全部「選択: A」で finalize した作業フォルダを返す。"""
    assert run("init", str(base), "--color", "navy")[0] == 0
    work = paths.resolve_work(base)
    ans = "---\n---\n"
    for i, n in enumerate(names, 1):
        (work / "slides" / f"{i:02d}-A.html").write_text(
            (pictures / f"{n}.html").read_text(encoding="utf-8"), encoding="utf-8"
        )
        ans += f"## S{i:02d}\n- 選択: A\n{answer_tail}"
    (work / "answers-1.md").write_text(ans, encoding="utf-8")
    assert run("finalize", str(base))[0] == 0
    return work


# ---- build ----


@slow
def test_見本はビルドでき_3倍で透過の画像として貼られ_題と注釈は文字のまま(tmp_path, run, make_work, sample):
    from PIL import Image
    from pptx import Presentation

    make_work(sample)
    base = make_work.base
    rc, out, err = run("build", str(base))
    assert rc == 0, out + err
    result = json.loads(out)
    assert result["ok"], result["violations"]
    assert result["violations"] == []
    prs = Presentation(str(base.parent.parent / "demo.pptx"))
    pics = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pics) == 1
    pic = pics[0]
    assert pic.name == "pp-picture"
    assert (pic.left / 9525, pic.top / 9525, pic.width / 9525, pic.height / 9525) == (64, 224, 1152, 366)
    assert "想定外" in pic._element.nvPicPr.cNvPr.get("descr")
    im = Image.open(io.BytesIO(pic.image.blob))
    assert im.size == (1152 * 3, 366 * 3)  # 3 倍の解像度
    assert im.mode == "RGBA"
    assert im.getpixel((0, 0))[3] == 0  # 透過（下地・ほかの部品が写り込まない）
    # 題・ヘッドメッセージ・注釈は、編集できる文字のまま
    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert any("通信経路" in t for t in texts)
    assert any("遮断が必要" in t for t in texts)


@slow
def test_全見本を1つの資料にビルドすると違反が0件で各スライドに絵が1つある(tmp_path, run, pictures):
    from pptx import Presentation

    base = tmp_path / "out" / "pptx-deck-source" / "demo"
    init_with_samples(run, base, pictures, SAMPLE_NAMES)
    rc, out, err = run("build", str(base))
    assert rc == 0, out + err
    assert json.loads(out)["violations"] == []
    prs = Presentation(str(base.parent.parent / "demo.pptx"))
    assert len(prs.slides) == 5
    for s in prs.slides:
        assert len([x for x in s.shapes if x.name == "pp-picture"]) == 1


@pytest.mark.parametrize("name", SAMPLE_NAMES)
def test_見本の先頭コメントに用途_差し替える所_守ること_data_altがある(pictures, name):
    head = (pictures / f"{name}.html").read_text(encoding="utf-8").split("-->", 1)[0]
    for word in ("用途", "差し替える所", "守ること", "data-alt"):
        assert word in head, f"{name}.html の先頭のコメントに「{word}」がありません"


@slow
def test_利用者の画像を絵の土台にして貼れる(tmp_path, run, pictures):
    """利用者の画像（source/material/）を土台にして、矢印・番号だけを重ねる。画像は絵の一部として貼られる。"""
    from PIL import Image
    from pptx import Presentation

    frag = (pictures / "annotate.html").read_text(encoding="utf-8")
    a = frag.index("<!-- 土台（画面の絵）")
    b = frag.index("<!-- 注釈 -->")
    frag = (
        frag[:a]
        + '<image href="material/screen.png" x="8" y="8" width="720" height="350" preserveAspectRatio="none"/>\n      '
        + frag[b:]
    )
    base = tmp_path / "out" / "pptx-deck-source" / "demo"
    assert run("init", str(base), "--color", "navy")[0] == 0
    work = paths.resolve_work(base)
    (work / "material").mkdir()
    Image.new("RGB", (720, 350), (200, 30, 30)).save(
        work / "material" / "screen.png"
    )  # 赤一色（テーマの色と区別できる）
    (work / "slides" / "01-A.html").write_text(frag, encoding="utf-8")
    (work / "answers-1.md").write_text("---\n---\n## S01\n- 選択: A\n- 判定: 承認\n- コメント: \n", encoding="utf-8")
    assert run("finalize", str(base))[0] == 0
    rc, out, err = run("check", str(work / "final.html"))
    assert (rc, json.loads(out)) == (0, []), out + err  # 画像の色は変数の検査の対象外
    rc, out, err = run("build", str(base))
    assert rc == 0, out + err
    pic = [s for s in Presentation(str(base.parent.parent / "demo.pptx")).slides[0].shapes if s.name == "pp-picture"][0]
    r, g, _b, alpha = (
        Image.open(io.BytesIO(pic.image.blob)).convert("RGBA").getpixel((3 * 100, 3 * 300))
    )  # 土台の画像の上（番号・凡例のない所）
    assert r > 150
    assert g < 80
    assert alpha == 255


# ---- svgids ----


def test_uniquifyは絵の中のidと参照だけに接頭辞を付ける():
    frag = (
        '<div class="pp-picture" data-alt="x"><svg><defs><marker id="m-acc"/><symbol id="i-server"/><symbol '
        'id="i-server2"/></defs>'
        '<use href="#i-server"/><use href="#i-server2"/><path marker-end="url(#m-acc)"/></svg></div><div id="m-acc">'
        "</div>"
    )
    out = svgids.uniquify(frag, "s01a1-")
    assert '<marker id="s01a1-m-acc"/>' in out
    assert '<use href="#s01a1-i-server"/>' in out
    assert '<use href="#s01a1-i-server2"/>' in out  # i-server と i-server2 を取り違えない
    assert 'marker-end="url(#s01a1-m-acc)"' in out
    assert out.endswith('<div id="m-acc"></div>')  # 絵の外は変えない


def test_uniquifyは絵のない断片を変えない():
    assert svgids.uniquify("<p>絵なし</p>", "p-") == "<p>絵なし</p>"


@slow
def test_あとの絵は前に同じid定義の絵があってもアイコンと矢じりが消えない(tmp_path, run, pictures):
    """絵だけを撮った PNG が、単独で作ったものと同じ。"""
    from PIL import Image, ImageChops
    from pptx import Presentation

    def picture_png(names, sub):
        base = tmp_path / sub / "out" / "pptx-deck-source" / "demo"
        init_with_samples(run, base, pictures, names, answer_tail="")
        rc, out, err = run("build", str(base))
        assert rc == 0, out + err
        last = list(Presentation(str(base.parent.parent / "demo.pptx")).slides)[-1]
        pic = [s for s in last.shapes if s.name == "pp-picture"][0]
        return Image.open(io.BytesIO(pic.image.blob)).convert("RGBA")

    alone = picture_png(["network"], "alone")
    after = picture_png(["traffic-flow", "network"], "after")
    assert alone.size == after.size
    diff = (
        ImageChops.difference(alone, after).convert("L").point(lambda v: 255 if v > 40 else 0)
    )  # 破線の端のアンチエイリアスの揺れは許す
    changed = diff.histogram()[255]
    assert changed / (alone.size[0] * alone.size[1]) < 0.002, (
        "2 枚目の絵が、単独で作ったものと違います（アイコン・矢じりが消えている）"
    )


# ---- 案が 1 つの流れ ----


@slow
def test_案が1つのとき選択だけの回答でfinalizeでき承認は最終確認に任せる(tmp_path, run, sample):
    base = tmp_path / "out" / "pptx-deck-source" / "demo"
    assert run("init", str(base), "--color", "navy")[0] == 0
    work = paths.resolve_work(base)
    (work / "slides" / "01-A.html").write_text(sample, encoding="utf-8")
    (work / "answers-1.md").write_text("---\n---\n## S01\n- 選択: A\n", encoding="utf-8")
    assert run("finalize", str(base))[0] == 0
    assert (work / "final-review.html").exists()
    assert run("build", str(base))[0] == 0
    assert (base.parent.parent / "demo.pptx").exists()


# ---- 検査 ----


def check_rules(run, work):
    rc, out, _ = run("check", str(work / "final.html"))
    return rc, {(v["rule"], v["severity"]) for v in json.loads(out)}


def drop_alt(s):
    return re.sub(r' data-alt="[^"]*"', "", s)


def drop_svg(s):
    return re.sub(r"<svg.*</svg>", "", s, flags=re.S)


def add_red_rect(s):
    return s.replace(
        '<text x="22" y="36"', '<rect x="1" y="1" width="9" height="9" fill="#ff0000"/><text x="22" y="36"'
    )


def add_var_rect(s):
    return s.replace(
        '<text x="22" y="36"', '<rect x="1" y="1" width="9" height="9" style="fill:var(--accent)"/><text x="22" y="36"'
    )


def shrink_font(s):
    return s.replace(
        '<text x="94" y="236" text-anchor="middle">', '<text x="94" y="236" text-anchor="middle" font-size="14">'
    )


def double_viewbox(s):
    # viewBox を倍にすると描画は半分（20 → 10px）になり、最小の文字の大きさで落ちる
    return s.replace('viewBox="0 0 1152 366"', 'viewBox="0 0 2304 732"')


def text_outside(s):
    return s.replace('<text x="22" y="36"', '<text x="1100" y="36"')


def overlapping_texts(s):
    return s.replace(
        '<text x="244" y="36" font-weight="700">DMZ</text>',
        '<text x="244" y="36" font-weight="700">DMZ</text><text x="250" y="38">DMZ2</text>',
    )


def overlap_other_part(s):
    return s.replace(
        "</section>",
        '<div class="pp-text" style="left:100px;top:300px;width:300px;height:40px">重なる文字</div>\n</section>',
    )


def overlap_inside(s):
    return s.replace(
        "</svg>",
        '<rect x="100" y="100" width="300" height="100" class="sv-f-light"/><rect x="150" y="120" width="300" '
        'height="100" class="sv-f-mid"/></svg>',
    )


@slow
def test_見本は検査に違反がない(run, make_work, sample):
    assert check_rules(run, make_work(sample)) == (0, set())


@slow
@pytest.mark.parametrize(
    ("mutate", "rc", "present", "absent"),
    [
        (drop_alt, 1, ("imageAlt", "error"), None),
        (drop_svg, 1, ("pictureSvg", "error"), None),
        (add_red_rect, 0, ("pictureColor", "warning"), None),  # 警告。ビルドは止めない
        (add_var_rect, None, None, ("pictureColor", "warning")),  # 変数・クラスの色は警告しない
        (shrink_font, 1, ("fontMin", "error"), None),
        (double_viewbox, None, ("fontMin", "error"), None),  # viewBox の倍率を加味する
        (text_outside, 1, ("pictureOverflow", "error"), None),
        (overlapping_texts, None, ("pictureTextOverlap", "warning"), None),
        (overlap_other_part, 1, ("overlap", "error"), None),  # 絵と他の部品の重なりは検査する
        (overlap_inside, None, None, ("overlap", "error")),  # 絵の内部の重なりは検査しない
    ],
    ids=[
        "altなし",
        "svgなし",
        "直書きの色",
        "変数の色",
        "小さい文字",
        "viewBox倍率",
        "絵の外へ出た文字",
        "文字の重なり",
        "他部品との重なり",
        "絵の内部の重なり",
    ],
)
def test_絵の検査(run, make_work, sample, mutate, rc, present, absent):
    got_rc, rules = check_rules(run, make_work(mutate(sample)))
    if rc is not None:
        assert got_rc == rc
    if present:
        assert present in rules
    if absent:
        assert absent not in rules


# ---- 部品の素材 ----


def test_parts_svgは正しいXMLで必要なidを持ち色はクラスだけで指定する(pkg):
    text = (pkg / "assets" / "pictures" / "parts.svg").read_text(encoding="utf-8")
    root = ET.fromstring(text)  # noqa: S314  # 同梱の自作 SVG を読むだけ（外部入力ではない）
    ids = {e.get("id") for e in root.iter() if e.get("id")}
    need = {
        "i-user",
        "i-pc",
        "i-server",
        "i-db",
        "i-router",
        "i-switch",
        "i-fw",
        "i-lb",
        "i-cloud",
        "i-internet",
        "i-alert",
        "i-doc",
        "m-acc",
        "m-emph",
    }
    assert need <= ids, need - ids
    assert re.findall(r'(?:fill|stroke)="(?!none)[^"]*"', text) == []  # 色はクラスだけ


def test_見本のsymbolは部品ライブラリと一致する(pkg, sample):
    parts = (pkg / "assets" / "pictures" / "parts.svg").read_text(encoding="utf-8")
    sids = re.findall(r'<symbol id="([^"]+)"', sample)
    assert sids, "見本に symbol がありません"
    for sid in sids:
        assert f'<symbol id="{sid}"' in parts
        a = re.search(rf'<symbol id="{sid}".*?</symbol>', sample, re.S).group(0)
        b = re.search(rf'<symbol id="{sid}".*?</symbol>', parts, re.S).group(0)
        assert re.sub(r"\s+", "", a) == re.sub(r"\s+", "", b), sid
