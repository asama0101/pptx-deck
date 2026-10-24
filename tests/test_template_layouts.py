"""テンプレートの全レイアウトから、章扉・目次・型の専用レイアウトを推定し、PPTX に使う。
（init の推定・template-map の修正・build）"""

# ruff: noqa: S101  # pytest の assert
import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from pptx_deck import cli, paths, template

NAMES = {0: "表紙", 1: "本文", 2: "章扉", 3: "目次", 4: "スケジュール", 5: "裏表紙"}


def make_template(path: Path, names: dict[int, str]) -> Path:
    """python-pptx の既定テンプレート（11 レイアウト）を 16:9 にして、レイアウト名を付け替えて保存する。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    for i, name in names.items():
        prs.slide_layouts[i]._element.cSld.set("name", name)
    for lay in prs.slide_layouts:  # 既定の配置は 4:3 用なので、本文のタイトルを標準テンプレートと同じ位置・幅にする
        for ph in lay.placeholders:
            if ph.placeholder_format.type in (1, 3):  # TITLE・CENTER_TITLE
                ph.left, ph.top, ph.width, ph.height = 64 * 9525, 36 * 9525, 1152 * 9525, 50 * 9525
    prs.save(str(path))
    return path


def strip_titles(path: Path, idxs) -> Path:
    prs = Presentation(str(path))
    for i in idxs:
        for ph in list(prs.slide_layouts[i].placeholders):
            if ph.placeholder_format.type == 1:  # TITLE
                ph._element.getparent().remove(ph._element)
    prs.save(str(path))
    return path


@pytest.fixture
def tpl(tmp_path):
    return make_template(tmp_path / "t.pptx", NAMES)


def run(capsys, *argv):
    rc = cli.main(list(argv))
    cap = capsys.readouterr()
    return rc, cap.out, cap.err


def build_deck(capsys, base, comp, frags):
    """init 済みの作業フォルダに断片を並べ、全スライドを A 案で承認して build し、PPTX のレイアウト名を返す。"""
    work = paths.resolve_work(base)
    for n, rel in enumerate(frags, 1):
        shutil.copyfile(comp / rel, work / "slides" / f"{n:02d}-A.html")
    assert run(capsys, "variants", str(base))[0] == 0
    (work / "answers-1.md").write_text(
        "---\ndeck: demo\nround: 1\n---\n"
        + "".join(f"## S{i:02d}\n- 選択: A\n- 判定: 承認\n- コメント: \n" for i in range(1, len(frags) + 1)),
        encoding="utf-8",
    )
    assert run(capsys, "finalize", str(base))[0] == 0
    rc, out, err = run(capsys, "build", str(base))
    assert rc == 0, out + err
    return [s.slide_layout.name for s in Presentation(str(base / "demo.pptx")).slides]


# ---- 推定 ----


def test_章扉と目次と型の専用レイアウトを名前で推定する(tpl):
    info = template.analyze(tpl)
    assert info.kinds == {"section": 2, "agenda": 3}
    assert {t: i for t, i in info.types.items() if i == 4} == {"gantt": 4, "roadmap": 4, "timeline": 4}
    assert info.closing == 5


def test_標準テンプレートは専用レイアウトを持たず本文で代える旨を伝える():
    info = template.analyze(paths.ASSETS / "template" / "standard.potx")
    assert (info.kinds, info.types) == ({}, {})
    assert any("章扉" in n for n in info.notes)


def test_英語の既定名のTwoContentは目次にしない(tmp_path):
    info = template.analyze(make_template(tmp_path / "en.pptx", {}))
    assert "agenda" not in info.kinds


def test_TableOfContentsは目次にする(tmp_path):
    info = template.analyze(make_template(tmp_path / "toc.pptx", {3: "Table of Contents"}))
    assert info.kinds.get("agenda") == 3


def test_表紙の派生レイアウトは表の型にしない(tmp_path):
    names = {0: "表紙", 1: "本文", 5: "裏表紙", 3: "表紙 2"}
    info = template.analyze(make_template(tmp_path / "cv.pptx", names))
    assert 3 not in info.types.values()


def test_to_dictとfrom_dictで往復し旧いtemplate_jsonも読める(tpl):
    info = template.analyze(tpl)
    again = template.TemplateInfo.from_dict(json.loads(json.dumps(info.to_dict())))
    assert again.kinds == {"section": 2, "agenda": 3}
    assert again.types["gantt"] == 4
    old = info.to_dict()
    for k in ("kinds", "types", "layout_keepout"):
        old.pop(k)
    assert template.TemplateInfo.from_dict(old).kinds == {}


@pytest.mark.parametrize(
    ("slide", "expected"),
    [
        ({"layout": "section", "type": "section"}, 2),
        ({"layout": "diagram", "type": "gantt"}, 4),  # 型の専用レイアウトが先
        ({"layout": "diagram", "type": "flow-h"}, 1),  # 専用がない型は本文
        ({"layout": "diagram"}, 1),  # 自由設計の案（型なし）
    ],
    ids=["章扉", "型の専用が先", "専用なしの型は本文", "型なしは本文"],
)
def test_layout_indexは型の専用レイアウトを優先する(tpl, slide, expected):
    info = template.analyze(tpl)
    assert template.layout_index(info, slide) == expected


def test_layout_indexは表紙に表紙を返す(tpl):
    info = template.analyze(tpl)
    assert template.layout_index(info, {"layout": "cover", "type": "cover"}) == info.cover == 0


def test_set_mappingは対応を足し外し不正はValueErrorにする(tpl):
    info = template.analyze(tpl)
    info = template.set_mapping(tpl, info, "agenda", "本文")
    assert info.kinds["agenda"] == 1
    info = template.set_mapping(tpl, info, "gantt", None)
    assert "gantt" not in info.types
    info = template.set_mapping(tpl, info, "flow-h", "スケジュール")
    assert info.types["flow-h"] == 4
    with pytest.raises(ValueError, match="nonexistent"):
        template.set_mapping(tpl, info, "nonexistent", "本文")
    with pytest.raises(ValueError, match="存在しないレイアウト"):
        template.set_mapping(tpl, info, "agenda", "存在しないレイアウト")


# ---- build ----


@pytest.mark.slow
def test_PPTXは推定した専用レイアウトで作られる(tmp_path, capsys, pkg, tpl):
    base = tmp_path / "demo"
    rc, out, err = run(capsys, "init", str(base), "--template", str(tpl))
    assert rc == 0, err
    assert "layout_map" in out
    names = build_deck(
        capsys,
        base,
        pkg / "assets" / "components",
        ["layouts/cover.html", "layouts/section.html", "layouts/agenda.html", "diagrams/gantt.html", "charts/bar.html"],
    )
    assert names == ["表紙", "章扉", "目次", "スケジュール", "本文"]


@pytest.mark.slow
def test_template_mapの変更がPPTXのレイアウトに反映される(tmp_path, capsys, pkg, tpl):
    base = tmp_path / "demo"
    assert run(capsys, "init", str(base), "--template", str(tpl))[0] == 0
    rc, out, err = run(capsys, "template-map", str(base), "--clear", "section", "--set", "bar=スケジュール")
    assert rc == 0, err
    cfg = json.loads((paths.resolve_work(base) / "template.json").read_text(encoding="utf-8"))
    assert "section" not in cfg["info"]["kinds"]
    assert cfg["info"]["types"]["bar"] == 4
    # 章扉の対応を外したので章扉は本文、棒グラフはスケジュールのレイアウトで作られる
    names = build_deck(capsys, base, pkg / "assets" / "components", ["layouts/section.html", "charts/bar.html"])
    assert names == ["本文", "スケジュール"]


def test_template_mapは決まらない指定を終了コード2で断る(tmp_path, capsys, tpl):
    base = tmp_path / "demo"
    assert run(capsys, "init", str(base), "--template", str(tpl))[0] == 0
    rc, _, err = run(capsys, "template-map", str(base), "--set", "bar=ない")
    assert rc == 2
    assert "決まりません" in err


# ---- 壊れたテンプレート ----


def test_本文レイアウトにタイトルがなければタイトルのあるレイアウトで代える(tmp_path):
    tpl = strip_titles(make_template(tmp_path / "t.pptx", NAMES), [1])
    info = template.analyze(tpl)
    assert info.body != 1


def test_タイトルのあるレイアウトが1つもなければ分かりやすいエラーで止まる(tmp_path, capsys):
    tpl = strip_titles(make_template(tmp_path / "t.pptx", NAMES), range(11))
    with pytest.raises(ValueError, match="本文に使えるレイアウト"):
        template.analyze(tpl)
    rc, _, err = run(capsys, "init", str(tmp_path / "d"), "--template", str(tpl))
    assert rc == 2
    assert "エラー" in err


@pytest.mark.parametrize("pptxでない", [False, True], ids=["存在しない", "pptxではない"])
def test_存在しないかpptxでないテンプレートは終了コード2で止まる(tmp_path, capsys, pptxでない):
    t = tmp_path / "nai.pptx"
    if pptxでない:
        t.write_text("pptx ではない", encoding="utf-8")
    rc, _, err = run(capsys, "init", str(tmp_path / "d2"), "--template", str(t))
    assert rc == 2
    assert "エラー" in err


@pytest.mark.parametrize("cmd", ["build", "draft", "finalize", "check"])
def test_template_jsonがなければ終了コード2で案内する(tmp_path, capsys, cmd):
    deck = tmp_path / "empty"
    (deck / "source").mkdir(parents=True)
    (deck / "source" / "final.html").write_text("<html></html>", encoding="utf-8")
    target = str(deck / "source" / "final.html") if cmd == "check" else str(deck)
    rc, _, err = run(capsys, cmd, target)
    assert rc == 2
    assert "template.json がありません" in err
