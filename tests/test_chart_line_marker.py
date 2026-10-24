"""折れ線のマーカーの塗りと枠が、系列の線の色になる。"""

# ruff: noqa: S101  # pytest の assert
from types import SimpleNamespace

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx_deck import charts

PAL = SimpleNamespace(accent="#00AA55", accent_mid="#55AA88", accent_dark="#005522", emphasis="#CC0000", text="#222222")


def _rgb(parent, *path):
    el = parent
    for tag in path:
        el = el.find(qn(tag))
        assert el is not None, tag
    return el.get("val")


def test_マーカーの塗りと枠は系列の線の色になる():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = {
        "type": "line",
        "categories": ["a", "b", "c"],
        "series": [{"name": "x", "values": [1, 2, 3]}, {"name": "y", "values": [3, 2, 1]}],
    }
    charts.add_chart(slide, {"chart": spec, "rect": {"x": 0, "y": 0, "w": 400, "h": 300}}, PAL)
    chart = [s for s in slide.shapes if s.has_chart][0].chart
    colors = []
    for ser in chart.plots[0].series:
        x = ser._element
        line = _rgb(x, "c:spPr", "a:ln", "a:solidFill", "a:srgbClr")
        sp = x.find(qn("c:marker")).find(qn("c:spPr"))
        assert _rgb(sp, "a:solidFill", "a:srgbClr") == line
        assert _rgb(sp, "a:ln", "a:solidFill", "a:srgbClr") == line
        colors.append(line)
    assert colors == ["00AA55", "55AA88"]
