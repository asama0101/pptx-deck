"""ウォーターフォールの小さい段のラベルが fmt（桁数・単位）に従う。"""

# ruff: noqa: S101  # pytest の assert
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx_deck import charts

PAL = SimpleNamespace(accent="#00AA55", accent_mid="#55AA88", accent_dark="#005522", emphasis="#CC0000", text="#222222")


@pytest.mark.parametrize(
    ("value", "fmt", "want"),
    [
        (0.5, "#,##0.0", "0.5"),
        (300, '#,##0"万円"', "300万円"),
        (1234, "#,##0", "1,234"),
        (5, "0", "5"),
        (2.5, "0", "3"),
        (-2.5, "0", "-3"),
        (0.125, "0.00", "0.13"),
        (0.05, "0.0%", "5.0%"),
        (0.5, "0%", "50%"),
        (0.125, "0%", "13%"),
        (1234.5, '[>=2000]#,##0.0;""', "1,234.5"),
        (0.25, "General", "0.25"),
    ],
    ids=[
        "小数1桁",
        "引用符つきの単位",
        "桁区切り",
        "整数",
        "半分は0から遠い側へ丸める",
        "負の半分も0から遠い側",
        "小数2桁の丸め",
        "パーセント小数1桁",
        "パーセント整数",
        "パーセントの丸め",
        "条件と区画は無視",
        "解釈できない書式は素の数字",
    ],
)
def test_ラベル文字は書式の最初の区画に従う(value, fmt, want):
    assert charts._label_text(value, fmt) == want


def _waterfall(spec_extra, values):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = {
        "type": "waterfall",
        "categories": ["a", "b", "c"],
        "series": [{"name": "x", "values": values}],
        **spec_extra,
    }
    charts.add_chart(slide, {"chart": spec, "rect": {"x": 0, "y": 0, "w": 400, "h": 300}}, PAL)
    return [s for s in slide.shapes if s.has_chart][0].chart


def test_小さい段のラベルは_fmt_の桁数に従う():
    chart = _waterfall({"fmt": "#,##0.0"}, [100, 0.5, 100.5])
    assert chart.plots[0].series[0].points[1].data_label.text_frame.text == "0.5"


def test_データラベルの書式は最初の区画だけを使う():
    chart = _waterfall({"fmt": "[>=1]#,##0;△#,##0"}, [100, 5, 105])
    assert chart.plots[0].series[1].data_labels.number_format == '[>=8.4]#,##0;""'
