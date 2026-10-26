"""行間は、フォントの行の高さに依存しない固定のポイントで書く（HTML の line-height × 文字サイズ）。"""

# ruff: noqa: S101  # pytest の assert
import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches
from pptx_deck import build


def _para(size_px, lh, mb=0):
    run = {"text": "あ", "size_px": size_px, "bold": False, "italic": False, "color": "000000"}
    return {"align": "left", "bullet": False, "lh": lh, "mb": mb, "runs": [run]}


def _lnspc_pts(paras):
    """段落を書き出し、各段落の a:lnSpc の中身を（要素名, 値）で返す。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2))
    build.fill_text_frame(tb.text_frame, paras)
    out = []
    for p in tb.text_frame.paragraphs:
        child = p._p.pPr.find(qn("a:lnSpc"))[0]
        out.append((child.tag.split("}")[1], child.get("val")))
    return out


@pytest.mark.parametrize(
    ("size_px", "lh", "expected"),
    [
        (20, 1.5, "2250"),  # 20px × 1.5 = 30px = 22.5pt
        (16, 1.2, "1440"),  # 16px × 1.2 = 19.2px = 14.4pt
        (28, 1.35, "2835"),  # 28px × 1.35 = 37.8px = 28.35pt
    ],
    ids=["本文 20px・1.5", "16px・1.2", "題 28px・1.35"],
)
def test_行間は固定ポイントで書く(size_px, lh, expected):
    assert _lnspc_pts([_para(size_px, lh)]) == [("spcPts", expected)]


def test_段落に大きさの違う文字が混ざるときは最大の文字サイズで行間を決める():
    para = _para(12, 1.5)
    para["runs"].append({**para["runs"][0], "size_px": 20})
    assert _lnspc_pts([para]) == [("spcPts", "2250")]


def test_段落ごとに自分の文字サイズで行間を決める():
    assert _lnspc_pts([_para(20, 1.5), _para(12, 1.5)]) == [("spcPts", "2250"), ("spcPts", "1350")]
